from __future__ import annotations

import os

import re
import math

from itertools import combinations
from typing import Mapping, Optional, Sequence

import numpy as np
import pandas as pd
import matplotlib.pyplot as plt

from scipy.stats import rankdata, wilcoxon
from statsmodels.stats.multitest import multipletests
from statsmodels.stats.contingency_tables import mcnemar



# ============================================================
# Global configuration
# ============================================================

METHODS = [
    "ProACTS",
    "DiCE4EL",
    "DiCE4EL+",
]

ID_COLUMNS = [
    "template_id",
    "case:concept:name",
]


# ============================================================
# Plot configuration
# ============================================================

# Two main method colors.
METHOD_COLORS = {
    "ProACTS": "#4C78A8",
    "DiCE4EL": "#F58518",
    "DiCE4EL+": "#F58518",
}

# Three outcome colors for propagation-penalty comparisons.
PROPAGATION_OUTCOME_COLORS = {
    "ProACTS better": "#4C78A8",
    "Tie": "#BAB0AC",
    "Baseline better": "#F58518",
}


# ============================================================
# Metric specifications
# ============================================================

# ------------------------------------------------------------
# Success conditions
#
# ProACTS:
#   autoregressive loss
#
# DiCE4EL / DiCE4EL+:
#   non-autoregressive loss
# ------------------------------------------------------------

SUCCESS_SPECS = {
    "Flip Success": {
        "ProACTS": "margin_loss_flipped",
        "DiCE4EL": "margin_loss_na_flipped",
        "DiCE4EL+": "margin_loss_na_flipped",
    },

    "Margin Success": {
        "ProACTS": "margin_loss",
        "DiCE4EL": "margin_loss_na",
        "DiCE4EL+": "margin_loss_na",
    },
}


# ------------------------------------------------------------
# Main unconditional results
# ------------------------------------------------------------

MAIN_METRIC_SPECS = {
    "Margin Loss": {
        "ProACTS": "margin_loss",
        "DiCE4EL": "margin_loss_na",
        "DiCE4EL+": "margin_loss_na",
    },

    "Flip Loss": {
        "ProACTS": "margin_loss_flipped",
        "DiCE4EL": "margin_loss_na_flipped",
        "DiCE4EL+": "margin_loss_na_flipped",
    },

    "Propagation Penalty": {
        "ProACTS": "process_violation",
        "DiCE4EL": "process_violation",
        "DiCE4EL+": "process_violation",
    },

    "Distance": {
        "ProACTS": "distance",
        "DiCE4EL": "distance",
        "DiCE4EL+": "distance",
    },

    "Sparsity": {
        "ProACTS": "sparsity",
        "DiCE4EL": "sparsity",
        "DiCE4EL+": "sparsity",
    },
}


# ------------------------------------------------------------
# Conditional comparisons
# ------------------------------------------------------------

CONDITIONAL_METRIC_SPECS = {
    "Flip Success": {
        "Margin Loss": {
            "ProACTS": "margin_loss",
            "DiCE4EL": "margin_loss_na",
            "DiCE4EL+": "margin_loss_na",
        },

        "Propagation Penalty": {
            "ProACTS": "process_violation",
            "DiCE4EL": "process_violation",
            "DiCE4EL+": "process_violation",
        },

        "Distance": {
            "ProACTS": "distance",
            "DiCE4EL": "distance",
            "DiCE4EL+": "distance",
        },
    },

    "Margin Success": {
        "Margin Loss": {
            "ProACTS": "margin_loss",
            "DiCE4EL": "margin_loss_na",
            "DiCE4EL+": "margin_loss_na",
        },

        "Propagation Penalty": {
            "ProACTS": "process_violation",
            "DiCE4EL": "process_violation",
            "DiCE4EL+": "process_violation",
        },

        "Distance": {
            "ProACTS": "distance",
            "DiCE4EL": "distance",
            "DiCE4EL+": "distance",
        },
    },
}


LOWER_IS_BETTER = {
    "Flip Loss": True,
    "Margin Loss": True,
    "Propagation Penalty": True,
    "Distance": True,
}


# ============================================================
# ProACTS multiple desired positions
# ============================================================

PROACTS_MULTIPLE_METRICS = {
    "Margin Loss": "margin_loss",
    "Propagation Penalty": "process_violation",
    "Distance": "distance",
    "Sparsity": "sparsity",
}

DESIRED_COUNT_COLUMN = "num_desired"


# ============================================================
# Loading one best CF per factual query
# ============================================================

def _load_best_counterfactuals(
    excel_path: str,
    sheet_name: str = "CF_Fitness_Analysis",
    fitness_column: str = "fitness",
) -> pd.DataFrame:
    """
    Select one generated counterfactual per factual query.

    Selection is based ONLY on minimum internal fitness.
    """

    df = pd.read_excel(
        excel_path,
        sheet_name=sheet_name,
    )

    required = [
        *ID_COLUMNS,
        fitness_column,
    ]

    missing = [
        c
        for c in required
        if c not in df.columns
    ]

    if missing:
        raise KeyError(
            f"Missing required columns in '{excel_path}': "
            f"{missing}"
        )

    df = df.copy()

    df[fitness_column] = pd.to_numeric(
        df[fitness_column],
        errors="coerce",
    )

    df = df.loc[
        np.isfinite(
            df[fitness_column]
        )
    ].copy()

    if df.empty:
        raise ValueError(
            f"No finite '{fitness_column}' values "
            f"in '{excel_path}'."
        )

    best_indices = (
        df.groupby(
            ID_COLUMNS,
            dropna=False,
        )[fitness_column]
        .idxmin()
    )

    result = (
        df.loc[best_indices]
        .reset_index(drop=True)
    )

    if result.duplicated(
        subset=ID_COLUMNS,
        keep=False,
    ).any():
        raise RuntimeError(
            "Duplicate factual queries remained after "
            "best-CF selection."
        )

    return result


# ============================================================
# Load one seed
# ============================================================

def evaluate_seed(
    proacts_path: str,
    dice_path: str,
    dice4el_plus_path: str,
    proacts_multiple_path: Optional[str] = None,
    *,
    sheet_name: str = "CF_Fitness_Analysis",
    fitness_column: str = "fitness",
) -> dict[str, pd.DataFrame]:

    proacts_df = _load_best_counterfactuals(
        proacts_path,
        sheet_name=sheet_name,
        fitness_column=fitness_column,
    )

    dice_df = _load_best_counterfactuals(
        dice_path,
        sheet_name=sheet_name,
        fitness_column=fitness_column,
    )

    dice4el_plus_df = _load_best_counterfactuals(
        dice4el_plus_path,
        sheet_name=sheet_name,
        fitness_column=fitness_column,
    )

    result = {
        "ProACTS": proacts_df,
        "DiCE4EL": dice_df,
        "DiCE4EL+": dice4el_plus_df,

        "plot": pd.concat(
            [
                proacts_df.assign(method="ProACTS"),
                dice_df.assign(method="DiCE4EL"),
                dice4el_plus_df.assign(
                    method="DiCE4EL+"
                ),
            ],
            ignore_index=True,
        ),
    }

    if proacts_multiple_path is not None:

        proacts_multiple_df = _load_best_counterfactuals(
            proacts_multiple_path,
            sheet_name=sheet_name,
            fitness_column=fitness_column,
        )

        result["ProACTS_multiple"] = proacts_multiple_df

        result["proacts_multiple_plot"] = (
            proacts_multiple_df.assign(
                method="ProACTS"
            )
        )

    return result


# ============================================================
# Combine methods across seeds
# ============================================================

def _combine_method_results_across_seeds(
    all_results: Mapping[
        object,
        Mapping[str, pd.DataFrame],
    ],
) -> pd.DataFrame:

    frames = []

    for seed, seed_results in all_results.items():

        for method in METHODS:

            if method not in seed_results:
                raise KeyError(
                    f"Seed '{seed}' does not contain "
                    f"'{method}'."
                )

            frame = (
                seed_results[
                    method
                ].copy()
            )

            frame.insert(
                0,
                "Seed",
                seed,
            )

            frame.insert(
                1,
                "Method",
                method,
            )

            frames.append(
                frame
            )

    if not frames:
        raise ValueError(
            "No seed results supplied."
        )

    return pd.concat(
        frames,
        ignore_index=True,
        sort=False,
    )


# ============================================================
# Combine ProACTS multiple desired positions
# ============================================================

def _combine_proacts_multiple_across_seeds(
    seed_results: Mapping,
) -> pd.DataFrame:

    frames = []

    for seed, result in seed_results.items():

        if "ProACTS_multiple" not in result:
            raise KeyError(
                f"Seed '{seed}' does not contain "
                "'ProACTS_multiple'."
            )

        frame = (
            result[
                "ProACTS_multiple"
            ].copy()
        )

        frame.insert(
            0,
            "Seed",
            seed,
        )

        frames.append(
            frame
        )

    if not frames:
        raise ValueError(
            "No ProACTS multiple-position results."
        )

    return pd.concat(
        frames,
        ignore_index=True,
        sort=False,
    )


# ============================================================
# Formatting helpers
# ============================================================

def _format_number(
    value: float,
    digits: int = 6,
) -> str:

    if not np.isfinite(value):
        return "NA"

    if value == 0:
        return "0"

    if abs(value) < 10 ** (-(digits - 1)):
        return f"{value:.3e}"

    return f"{value:.{digits}g}"


def _format_p_value(
    value: float,
) -> str:

    if not np.isfinite(value):
        return "NA"

    if value < 0.001:
        return "<0.001"

    return f"{value:.3f}"


def _mean_sd_values(
    values: np.ndarray,
) -> tuple[float, float]:

    values = np.asarray(
        values,
        dtype=float,
    )

    mean = float(
        np.mean(values)
    )

    sd = (
        float(
            np.std(
                values,
                ddof=1,
            )
        )
        if len(values) > 1
        else np.nan
    )

    return mean, sd


def _mean_sd_text(
    values: np.ndarray,
) -> str:

    mean, sd = _mean_sd_values(
        values
    )

    if np.isfinite(sd):
        return (
            f"{_format_number(mean)} ± "
            f"{_format_number(sd)}"
        )

    return _format_number(
        mean
    )


def _median_iqr_values(
    values: np.ndarray,
) -> tuple[
    float,
    float,
    float,
]:

    q1, median, q3 = np.quantile(
        values,
        [
            0.25,
            0.50,
            0.75,
        ],
    )

    return (
        float(median),
        float(q1),
        float(q3),
    )


def _median_iqr_text(
    values: np.ndarray,
) -> str:

    median, q1, q3 = (
        _median_iqr_values(
            values
        )
    )

    return (
        f"{_format_number(median)} "
        f"[{_format_number(q1)}, "
        f"{_format_number(q3)}]"
    )


# ============================================================
# Aggregate metric per query across seeds
# ============================================================

def _aggregate_method_metric_across_seeds(
    combined: pd.DataFrame,
    method: str,
    source_column: str,
) -> pd.DataFrame:

    required = [
        "Seed",
        "Method",
        *ID_COLUMNS,
        source_column,
    ]

    missing = [
        c
        for c in required
        if c not in combined.columns
    ]

    if missing:
        raise KeyError(
            f"{method}/{source_column} missing: "
            f"{missing}"
        )

    df = combined.loc[
        combined["Method"].eq(method),
        [
            "Seed",
            *ID_COLUMNS,
            source_column,
        ],
    ].copy()

    df[source_column] = pd.to_numeric(
        df[source_column],
        errors="coerce",
    )

    df = df.loc[
        np.isfinite(
            df[source_column]
        )
    ]

    if df.empty:
        raise ValueError(
            f"No finite observations for "
            f"{method}/{source_column}."
        )

    return (
        df.groupby(
            ID_COLUMNS,
            dropna=False,
            as_index=False,
        )
        .agg(
            Value=(
                source_column,
                "median",
            ),
            N_Seeds=(
                "Seed",
                "nunique",
            ),
        )
    )


# ============================================================
# Pair two methods
# ============================================================

def _pair_metric(
    combined: pd.DataFrame,
    method_a: str,
    method_b: str,
    column_a: str,
    column_b: str,
) -> pd.DataFrame:

    a = (
        _aggregate_method_metric_across_seeds(
            combined,
            method_a,
            column_a,
        )
        .rename(
            columns={
                "Value": "A",
                "N_Seeds": "A N Seeds",
            }
        )
    )

    b = (
        _aggregate_method_metric_across_seeds(
            combined,
            method_b,
            column_b,
        )
        .rename(
            columns={
                "Value": "B",
                "N_Seeds": "B N Seeds",
            }
        )
    )

    paired = a.merge(
        b,
        on=ID_COLUMNS,
        how="inner",
        validate="one_to_one",
    )

    finite = (
        np.isfinite(
            paired["A"]
        )
        &
        np.isfinite(
            paired["B"]
        )
    )

    return (
        paired.loc[
            finite
        ]
        .reset_index(drop=True)
    )


# ============================================================
# Main results
# ============================================================

def _build_main_results_table(
    combined: pd.DataFrame,
) -> pd.DataFrame:

    rows = []

    for (
        metric_name,
        method_columns,
    ) in MAIN_METRIC_SPECS.items():

        method_frames = []

        for method in METHODS:

            frame = (
                _aggregate_method_metric_across_seeds(
                    combined,
                    method,
                    method_columns[
                        method
                    ],
                )
                .rename(
                    columns={
                        "Value": method,
                    }
                )
            )

            frame = frame[
                [
                    *ID_COLUMNS,
                    method,
                ]
            ]

            method_frames.append(
                frame
            )

        paired = method_frames[0]

        for frame in method_frames[1:]:

            paired = paired.merge(
                frame,
                on=ID_COLUMNS,
                how="inner",
                validate="one_to_one",
            )

        row = {
            "Metric": metric_name,
            "N Paired Instances": int(
                len(paired)
            ),
        }

        for method in METHODS:

            values = (
                paired[
                    method
                ]
                .to_numpy(
                    dtype=float
                )
            )

            row[
                f"{method} Mean ± SD"
            ] = _mean_sd_text(
                values
            )

            row[
                f"{method} Median [Q1, Q3]"
            ] = _median_iqr_text(
                values
            )

        rows.append(
            row
        )

    return pd.DataFrame(
        rows
    )


# ============================================================
# Success rates
# ============================================================

def _build_success_rate_table(
    combined: pd.DataFrame,
    tolerance: float,
) -> pd.DataFrame:

    raw_rows = []

    for (
        condition,
        method_columns,
    ) in SUCCESS_SPECS.items():

        for method in METHODS:

            source_column = (
                method_columns[
                    method
                ]
            )

            subset = combined.loc[
                combined[
                    "Method"
                ].eq(method),
                [
                    "Seed",
                    source_column,
                ],
            ].copy()

            subset[
                source_column
            ] = pd.to_numeric(
                subset[source_column],
                errors="coerce",
            )

            subset = subset.loc[
                np.isfinite(
                    subset[
                        source_column
                    ]
                )
            ]

            for (
                seed,
                seed_df,
            ) in subset.groupby(
                "Seed",
                sort=True,
            ):

                values = (
                    seed_df[
                        source_column
                    ]
                    .to_numpy(
                        dtype=float
                    )
                )

                success = (
                    values <= tolerance
                )

                raw_rows.append(
                    {
                        "Condition": condition,
                        "Method": method,
                        "Seed": seed,
                        "N Instances": len(
                            values
                        ),
                        "Success Count": int(
                            success.sum()
                        ),
                        "Success Rate": float(
                            success.mean()
                        ),
                    }
                )

    raw = pd.DataFrame(
        raw_rows
    )

    rows = []

    for (
        condition,
        method,
    ), group in raw.groupby(
        [
            "Condition",
            "Method",
        ],
        sort=False,
    ):

        rates = (
            group[
                "Success Rate"
            ]
            .to_numpy(
                dtype=float
            )
        )

        rows.append(
            {
                "Condition": condition,
                "Method": method,
                "N Seeds": len(
                    rates
                ),
                "Mean ± SD": (
                    _mean_sd_text(
                        rates
                    )
                ),
                "Mean Success Rate": float(
                    np.mean(
                        rates
                    )
                ),
            }
        )

    return pd.DataFrame(
        rows
    )



# ============================================================
# Flip / margin success significance
# ============================================================

def _build_success_significance_table(
    combined: pd.DataFrame,
    *,
    alpha: float = 0.001,
    tolerance: float = 1e-8,
) -> pd.DataFrame:
    """
    Statistically compare Flip Success and Margin Success.

    The experimental unit is the factual instance.

    For each success condition:
        1. Aggregate the corresponding loss per factual instance
           across seeds using the median.
        2. Convert the per-instance median loss to binary success:
               success = loss <= tolerance
        3. Compare ProACTS with each baseline using the exact
           paired McNemar test.
        4. Apply Holm correction across the two planned
           ProACTS-vs-baseline comparisons separately for each
           success condition.

    Planned comparisons:
        ProACTS vs DiCE4EL
        ProACTS vs DiCE4EL+

    Discordant pairs:
        ProACTS Only Success:
            ProACTS succeeds and the baseline fails.

        Baseline Only Success:
            the baseline succeeds and ProACTS fails.

    McNemar tests whether the two discordant counts differ.
    """

    comparisons = [
        ("ProACTS", "DiCE4EL"),
        ("ProACTS", "DiCE4EL+"),
    ]

    rows = []

    for condition, method_columns in SUCCESS_SPECS.items():

        for method_a, method_b in comparisons:

            # ------------------------------------------------
            # Pair the per-instance median success losses.
            #
            # _pair_metric() already aggregates each factual
            # instance across seeds using the median.
            # ------------------------------------------------

            paired = _pair_metric(
                combined=combined,
                method_a=method_a,
                method_b=method_b,
                column_a=method_columns[method_a],
                column_b=method_columns[method_b],
            )

            comparison = (
                f"{method_a} vs {method_b}"
            )

            if paired.empty:

                rows.append(
                    {
                        "Condition": condition,
                        "Comparison": comparison,
                        "Baseline": method_b,
                        "N Paired": 0,

                        "ProACTS Success Count": 0,
                        "Baseline Success Count": 0,

                        "ProACTS Success Rate": np.nan,
                        "Baseline Success Rate": np.nan,

                        "Both Success": 0,
                        "ProACTS Only Success": 0,
                        "Baseline Only Success": 0,
                        "Neither Success": 0,

                        "Discordant Pairs": 0,

                        "McNemar Raw p": np.nan,
                        "Direction": "Undefined",
                    }
                )

                continue

            # ------------------------------------------------
            # Convert per-instance median losses to success.
            # ------------------------------------------------

            success_a = (
                paired["A"]
                .to_numpy(dtype=float)
                <= tolerance
            )

            success_b = (
                paired["B"]
                .to_numpy(dtype=float)
                <= tolerance
            )

            n = int(
                len(success_a)
            )

            # ------------------------------------------------
            # Paired 2 x 2 contingency counts.
            #
            #                       Baseline
            #                   success   fail
            # ProACTS success     n11      n10
            #         fail        n01      n00
            #
            # McNemar uses n10 and n01.
            # ------------------------------------------------

            both_success = int(
                (
                    success_a
                    &
                    success_b
                ).sum()
            )

            proacts_only = int(
                (
                    success_a
                    &
                    ~success_b
                ).sum()
            )

            baseline_only = int(
                (
                    ~success_a
                    &
                    success_b
                ).sum()
            )

            neither_success = int(
                (
                    ~success_a
                    &
                    ~success_b
                ).sum()
            )

            discordant = (
                proacts_only
                +
                baseline_only
            )

            # ------------------------------------------------
            # Exact McNemar test.
            # ------------------------------------------------

            if discordant == 0:

                # The methods have the same success outcome for
                # every paired factual instance.
                raw_p = 1.0

            else:

                contingency_table = [
                    [
                        both_success,
                        proacts_only,
                    ],
                    [
                        baseline_only,
                        neither_success,
                    ],
                ]

                test_result = mcnemar(
                    contingency_table,
                    exact=True,
                )

                raw_p = float(
                    test_result.pvalue
                )

            # ------------------------------------------------
            # Direction from discordant pairs.
            # ------------------------------------------------

            if proacts_only > baseline_only:
                direction = "ProACTS"

            elif baseline_only > proacts_only:
                direction = method_b

            else:
                direction = "Tie"

            rows.append(
                {
                    "Condition": condition,
                    "Comparison": comparison,
                    "Baseline": method_b,
                    "N Paired": n,

                    "ProACTS Success Count": int(
                        success_a.sum()
                    ),
                    "Baseline Success Count": int(
                        success_b.sum()
                    ),

                    "ProACTS Success Rate": float(
                        success_a.mean()
                    ),
                    "Baseline Success Rate": float(
                        success_b.mean()
                    ),

                    "Both Success": both_success,
                    "ProACTS Only Success": proacts_only,
                    "Baseline Only Success": baseline_only,
                    "Neither Success": neither_success,

                    "Discordant Pairs": discordant,

                    "McNemar Raw p": raw_p,
                    "Direction": direction,
                }
            )

    results = pd.DataFrame(
        rows
    )

    # ========================================================
    # Holm correction
    #
    # Two planned tests are corrected separately within:
    #     Flip Success
    #     Margin Success
    # ========================================================

    results[
        "McNemar Holm-adjusted p"
    ] = np.nan

    for condition, group in results.groupby(
        "Condition",
        sort=False,
    ):

        valid = group.index[
            np.isfinite(
                group[
                    "McNemar Raw p"
                ]
            )
        ]

        if len(valid) == 0:
            continue

        _, adjusted, _, _ = multipletests(
            results.loc[
                valid,
                "McNemar Raw p",
            ].to_numpy(dtype=float),
            alpha=alpha,
            method="holm",
        )

        results.loc[
            valid,
            "McNemar Holm-adjusted p",
        ] = adjusted

    # ========================================================
    # Formatted values
    # ========================================================

    results[
        "Holm p"
    ] = (
        results[
            "McNemar Holm-adjusted p"
        ]
        .apply(
            _format_p_value
        )
    )

    results[
        "ProACTS Success %"
    ] = (
        100.0
        *
        results[
            "ProACTS Success Rate"
        ]
    )

    results[
        "Baseline Success %"
    ] = (
        100.0
        *
        results[
            "Baseline Success Rate"
        ]
    )

    # ========================================================
    # Conclusion
    # ========================================================

    conclusions = []

    for _, row in results.iterrows():

        p = row[
            "McNemar Holm-adjusted p"
        ]

        direction = row[
            "Direction"
        ]

        statistically_significant = (
            np.isfinite(p)
            and
            p < alpha
        )

        if (
            statistically_significant
            and
            direction
            not in {
                "Tie",
                "Undefined",
            }
        ):

            conclusion = (
                f"{direction} higher success rate"
            )

        else:

            conclusion = (
                "No significant difference"
            )

        conclusions.append(
            conclusion
        )

    results[
        "Conclusion"
    ] = conclusions

    return results[
        [
            "Condition",
            "Comparison",
            "Baseline",
            "N Paired",

            "ProACTS Success Count",
            "Baseline Success Count",

            "ProACTS Success Rate",
            "Baseline Success Rate",

            "ProACTS Success %",
            "Baseline Success %",

            "Both Success",
            "ProACTS Only Success",
            "Baseline Only Success",
            "Neither Success",
            "Discordant Pairs",

            "McNemar Raw p",
            "McNemar Holm-adjusted p",
            "Holm p",

            "Direction",
            "Conclusion",
        ]
    ].copy()



# ============================================================
# Seed robustness - single desired position
# ============================================================

def _build_seed_robustness_table(
    combined: pd.DataFrame,
) -> pd.DataFrame:

    rows = []

    for (
        metric_name,
        method_columns,
    ) in MAIN_METRIC_SPECS.items():

        for method in METHODS:

            source_column = (
                method_columns[
                    method
                ]
            )

            subset = combined.loc[
                combined[
                    "Method"
                ].eq(method),
                [
                    "Seed",
                    source_column,
                ],
            ].copy()

            subset[
                source_column
            ] = pd.to_numeric(
                subset[source_column],
                errors="coerce",
            )

            subset = subset.loc[
                np.isfinite(
                    subset[
                        source_column
                    ]
                )
            ]

            seed_means = []
            seed_medians = []

            for _, group in subset.groupby(
                "Seed",
                sort=True,
            ):

                values = (
                    group[
                        source_column
                    ]
                    .to_numpy(
                        dtype=float
                    )
                )

                seed_means.append(
                    np.mean(
                        values
                    )
                )

                seed_medians.append(
                    np.median(
                        values
                    )
                )

            rows.append(
                {
                    "Metric": metric_name,
                    "Method": method,
                    "N Seeds": len(
                        seed_means
                    ),
                    "Seed Mean ± SD": (
                        _mean_sd_text(
                            np.asarray(
                                seed_means
                            )
                        )
                    ),
                    "Seed Median ± SD": (
                        _mean_sd_text(
                            np.asarray(
                                seed_medians
                            )
                        )
                    ),
                }
            )

    return pd.DataFrame(
        rows
    )


# ============================================================
# Wilcoxon signed-rank variants
# ============================================================

def _run_signed_rank_test(
    values_a: np.ndarray,
    values_b: np.ndarray,
    *,
    zero_method: str,
    zero_tolerance: float = 1e-12,
) -> tuple[float, float]:
    """
    Run the paired Wilcoxon signed-rank test with a specified
    treatment of zero differences.

    zero_method="wilcox":
        Zero paired differences are removed before ranking.

    zero_method="pratt":
        Zero paired differences are included when assigning
        ranks, but their ranks do not contribute to either
        signed-rank sum.

    zero_method="zsplit":
        Zero paired differences are included when assigning
        ranks and their ranks are split equally between the
        positive and negative signed-rank sums.
    """

    if zero_method not in {
        "wilcox",
        "pratt",
        "zsplit",
    }:
        raise ValueError(
            "zero_method must be "
            "'wilcox', 'pratt', or 'zsplit'."
        )

    a = np.asarray(
        values_a,
        dtype=float,
    )

    b = np.asarray(
        values_b,
        dtype=float,
    )

    finite = (
        np.isfinite(a)
        &
        np.isfinite(b)
    )

    a = a[
        finite
    ]

    b = b[
        finite
    ]

    if len(a) == 0:
        return (
            np.nan,
            np.nan,
        )

    differences = (
        a
        -
        b
    )

    tied = np.isclose(
        differences,
        0.0,
        atol=zero_tolerance,
        rtol=0.0,
    )

    # Make our numerical tie definition explicit to SciPy.
    #
    # If two values differ only by a tiny numerical amount
    # within zero_tolerance, force their difference to zero.
    b = b.copy()

    b[
        tied
    ] = a[
        tied
    ]

    nonzero = (
        ~tied
    )

    # --------------------------------------------------------
    # All pairs tied
    # --------------------------------------------------------
    #
    # There is no evidence of a difference.
    #
    # For all three conventions we report:
    #     statistic = 0
    #     p = 1
    #
    # This also avoids edge-case behaviour in SciPy.
    # --------------------------------------------------------

    if not np.any(
        nonzero
    ):
        return (
            0.0,
            1.0,
        )

    try:

        result = wilcoxon(
            a,
            b,
            alternative="two-sided",
            zero_method=zero_method,
            method="auto",
        )

    except ValueError:

        return (
            np.nan,
            np.nan,
        )

    return (
        float(
            result.statistic
        ),
        float(
            result.pvalue
        ),
    )


def _run_wilcoxon(
    values_a: np.ndarray,
    values_b: np.ndarray,
    zero_tolerance: float = 1e-12,
) -> tuple[float, float]:

    return _run_signed_rank_test(
        values_a,
        values_b,
        zero_method="wilcox",
        zero_tolerance=zero_tolerance,
    )


def _run_pratt(
    values_a: np.ndarray,
    values_b: np.ndarray,
    zero_tolerance: float = 1e-12,
) -> tuple[float, float]:

    return _run_signed_rank_test(
        values_a,
        values_b,
        zero_method="pratt",
        zero_tolerance=zero_tolerance,
    )


def _run_zsplit(
    values_a: np.ndarray,
    values_b: np.ndarray,
    zero_tolerance: float = 1e-12,
) -> tuple[float, float]:

    return _run_signed_rank_test(
        values_a,
        values_b,
        zero_method="zsplit",
        zero_tolerance=zero_tolerance,
    )


# ============================================================
# Rank-biserial
# ============================================================

def _rank_biserial_correlation(
    values_a: np.ndarray,
    values_b: np.ndarray,
    zero_tolerance: float = 1e-12,
) -> float:

    differences = (
        np.asarray(
            values_a,
            dtype=float,
        )
        -
        np.asarray(
            values_b,
            dtype=float,
        )
    )

    keep = (
        np.isfinite(
            differences
        )
        &
        ~np.isclose(
            differences,
            0.0,
            atol=zero_tolerance,
            rtol=0.0,
        )
    )

    differences = (
        differences[
            keep
        ]
    )

    if len(
        differences
    ) == 0:
        return 0.0

    ranks = rankdata(
        np.abs(
            differences
        ),
        method="average",
    )

    positive = float(
        ranks[
            differences > 0
        ].sum()
    )

    negative = float(
        ranks[
            differences < 0
        ].sum()
    )

    total = (
        positive
        +
        negative
    )

    if total == 0:
        return 0.0

    return (
        positive
        -
        negative
    ) / total


def _effect_magnitude(
    effect: float,
) -> str:

    if not np.isfinite(
        effect
    ):
        return "Undefined"

    effect = abs(
        effect
    )

    if effect < 0.10:
        return "Negligible"

    if effect < 0.30:
        return "Small"

    if effect < 0.50:
        return "Medium"

    return "Large"


def _better_method_from_effect(
    method_a: str,
    method_b: str,
    effect_size: float,
    lower_is_better: bool,
    tolerance: float = 1e-12,
) -> str:

    if not np.isfinite(
        effect_size
    ):
        return "Undefined"

    if np.isclose(
        effect_size,
        0.0,
        atol=tolerance,
        rtol=0.0,
    ):
        return "Tie"

    if lower_is_better:

        return (
            method_a
            if effect_size < 0
            else method_b
        )

    return (
        method_a
        if effect_size > 0
        else method_b
    )


# ============================================================
# Common-success query subset
# ============================================================

def _get_common_success_queries(
    combined: pd.DataFrame,
    condition: str,
    method_a: str,
    method_b: str,
    tolerance: float,
) -> tuple[
    pd.DataFrame,
    dict,
]:

    columns = (
        SUCCESS_SPECS[
            condition
        ]
    )

    paired = _pair_metric(
        combined,
        method_a,
        method_b,
        columns[
            method_a
        ],
        columns[
            method_b
        ],
    )

    success_a = (
        paired["A"]
        .to_numpy(
            dtype=float
        )
        <= tolerance
    )

    success_b = (
        paired["B"]
        .to_numpy(
            dtype=float
        )
        <= tolerance
    )

    common = (
        success_a
        &
        success_b
    )

    retained = (
        paired.loc[
            common,
            ID_COLUMNS,
        ]
        .drop_duplicates()
        .reset_index(
            drop=True
        )
    )

    n_total = len(
        paired
    )

    n_retained = len(
        retained
    )

    return (
        retained,
        {
            "Condition": condition,
            "Comparison": (
                f"{method_a} vs "
                f"{method_b}"
            ),
            "N Paired": n_total,
            f"{method_a} Successful": int(
                success_a.sum()
            ),
            f"{method_b} Successful": int(
                success_b.sum()
            ),
            "N Retained Both Successful": (
                n_retained
            ),
            "Retained %": (
                n_retained
                / n_total
                if n_total
                else np.nan
            ),
        },
    )


# ============================================================
# Retained counts
# ============================================================

def _build_retained_counts_table(
    combined: pd.DataFrame,
    tolerance: float,
) -> pd.DataFrame:

    rows = []

    for condition in SUCCESS_SPECS:

        for (
            method_a,
            method_b,
        ) in combinations(
            METHODS,
            2,
        ):

            _, counts = (
                _get_common_success_queries(
                    combined,
                    condition,
                    method_a,
                    method_b,
                    tolerance,
                )
            )

            rows.append(
                counts
            )

    return pd.DataFrame(
        rows
    )


# ============================================================
# Propagation penalty classification
# ============================================================

def _classify_propagation_penalty_pairs(
    values_a: np.ndarray,
    values_b: np.ndarray,
    *,
    tolerance: float = 1e-12,
) -> pd.DataFrame:
    """
    A = ProACTS
    B = baseline
    """

    a = np.asarray(
        values_a,
        dtype=float,
    )

    b = np.asarray(
        values_b,
        dtype=float,
    )

    if len(a) != len(b):
        raise ValueError(
            "Propagation-penalty pair lengths differ."
        )

    equal = np.isclose(
        a,
        b,
        atol=tolerance,
        rtol=0.0,
    )

    both_zero = (
        np.isclose(
            a,
            0.0,
            atol=tolerance,
            rtol=0.0,
        )
        &
        np.isclose(
            b,
            0.0,
            atol=tolerance,
            rtol=0.0,
        )
    )

    outcome = np.full(
        len(a),
        "",
        dtype=object,
    )

    outcome[
        a < b - tolerance
    ] = "ProACTS better"

    outcome[
        a > b + tolerance
    ] = "Baseline better"

    outcome[
        equal
    ] = "Tie"

    tie_type = np.full(
        len(a),
        "",
        dtype=object,
    )

    tie_type[
        both_zero
    ] = "Both zero"

    tie_type[
        equal
        &
        ~both_zero
    ] = "Equal non-zero"

    return pd.DataFrame(
        {
            "Outcome": outcome,
            "Tie Type": tie_type,
        }
    )


# ============================================================
# Tie diagnostics
# ============================================================

def _build_propagation_penalty_tie_table(
    combined: pd.DataFrame,
    *,
    condition: str,
    baseline_method: str,
    tolerance: float,
) -> pd.DataFrame:

    retained, _ = (
        _get_common_success_queries(
            combined,
            condition,
            "ProACTS",
            baseline_method,
            tolerance,
        )
    )

    penalty_pairs = _pair_metric(
        combined,
        "ProACTS",
        baseline_method,
        "process_violation",
        "process_violation",
    )

    penalty_pairs = retained.merge(
        penalty_pairs,
        on=ID_COLUMNS,
        how="inner",
        validate="one_to_one",
    )

    classes = (
        _classify_propagation_penalty_pairs(
            penalty_pairs["A"],
            penalty_pairs["B"],
            tolerance=tolerance,
        )
    )

    n = len(classes)

    rows = []

    for category in [
        "ProACTS better",
        "Tie",
        "Baseline better",
    ]:

        count = int(
            (
                classes["Outcome"]
                == category
            ).sum()
        )

        rows.append(
            {
                "Condition": condition,
                "Comparison": (
                    f"ProACTS vs {baseline_method}"
                ),
                "Category": category,
                "Count": count,
                "Percent": (
                    count / n
                    if n
                    else np.nan
                ),
                "N Retained": n,
            }
        )

    for category in [
        "Both zero",
        "Equal non-zero",
    ]:

        count = int(
            (
                classes["Tie Type"]
                == category
            ).sum()
        )

        rows.append(
            {
                "Condition": condition,
                "Comparison": (
                    f"ProACTS vs {baseline_method}"
                ),
                "Category": category,
                "Count": count,
                "Percent": (
                    count / n
                    if n
                    else np.nan
                ),
                "N Retained": n,
            }
        )

    return pd.DataFrame(
        rows
    )


# ============================================================
# Conditional statistical comparisons
# ============================================================

def _build_conditional_pairwise_results(
    combined: pd.DataFrame,
    *,
    alpha: float,
    tolerance: float,
    practical_effect_threshold: float,
) -> pd.DataFrame:
    """
    Pairwise comparisons after retaining only queries where both
    methods satisfy the relevant success condition.

    Three Wilcoxon signed-rank zero-handling variants are
    reported:

        Wilcox:
            zero_method="wilcox"

            Zero paired differences are removed before ranking.

        Pratt:
            zero_method="pratt"

            Zero paired differences remain in rank assignment,
            but their ranks do not contribute to either signed
            rank sum.

        Z-split:
            zero_method="zsplit"

            Zero paired differences remain in rank assignment
            and their ranks are split equally between the
            positive and negative signed-rank sums.

    Holm correction is applied independently to each
    zero-handling variant within each (Condition, Metric)
    family.

    Rank-biserial correlation remains calculated on non-tied
    pairs and should therefore be interpreted together with
    Tie %.
    """

    rows = []

    # ========================================================
    # Conditions and metrics
    # ========================================================

    for (
        condition,
        metric_specs,
    ) in CONDITIONAL_METRIC_SPECS.items():

        for (
            metric_name,
            method_columns,
        ) in metric_specs.items():

            for (
                method_a,
                method_b,
            ) in combinations(
                METHODS,
                2,
            ):

                # =================================================
                # Retain queries where BOTH methods satisfy
                # the success condition.
                # =================================================

                (
                    retained_ids,
                    counts,
                ) = (
                    _get_common_success_queries(
                        combined,
                        condition,
                        method_a,
                        method_b,
                        tolerance,
                    )
                )

                base_row = {
                    "Condition": condition,

                    "Metric": metric_name,

                    "Comparison": (
                        f"{method_a} vs "
                        f"{method_b}"
                    ),

                    "N Paired Before Condition": (
                        counts[
                            "N Paired"
                        ]
                    ),
                }

                # =================================================
                # Nothing retained
                # =================================================

                if retained_ids.empty:

                    rows.append(
                        {
                            **base_row,

                            "N Retained": 0,
                            "Retained %": 0.0,

                            "N Non-Ties": 0,
                            "Tie %": np.nan,

                            "A Mean ± SD": "NA",
                            "B Mean ± SD": "NA",

                            "A Median [Q1, Q3]": "NA",
                            "B Median [Q1, Q3]": "NA",

                            "Median Difference A-B": (
                                np.nan
                            ),

                            "Wilcox Statistic": np.nan,
                            "Wilcox Raw p": np.nan,

                            "Pratt Statistic": np.nan,
                            "Pratt Raw p": np.nan,

                            "Z-split Statistic": np.nan,
                            "Z-split Raw p": np.nan,

                            "Rank-Biserial (Non-Ties)": (
                                np.nan
                            ),

                            "Effect": "Undefined",

                            "Better Method": (
                                "Undefined"
                            ),
                        }
                    )

                    continue

                # =================================================
                # Pair the actual metric
                # =================================================

                paired_metric = _pair_metric(
                    combined,
                    method_a,
                    method_b,
                    method_columns[
                        method_a
                    ],
                    method_columns[
                        method_b
                    ],
                )

                conditional = (
                    retained_ids.merge(
                        paired_metric,
                        on=ID_COLUMNS,
                        how="inner",
                        validate="one_to_one",
                    )
                )

                # =================================================
                # No metric pairs after merge
                # =================================================

                if conditional.empty:

                    rows.append(
                        {
                            **base_row,

                            "N Retained": 0,
                            "Retained %": 0.0,

                            "N Non-Ties": 0,
                            "Tie %": np.nan,

                            "A Mean ± SD": "NA",
                            "B Mean ± SD": "NA",

                            "A Median [Q1, Q3]": "NA",
                            "B Median [Q1, Q3]": "NA",

                            "Median Difference A-B": (
                                np.nan
                            ),

                            "Wilcox Statistic": np.nan,
                            "Wilcox Raw p": np.nan,

                            "Pratt Statistic": np.nan,
                            "Pratt Raw p": np.nan,

                            "Z-split Statistic": np.nan,
                            "Z-split Raw p": np.nan,

                            "Rank-Biserial (Non-Ties)": (
                                np.nan
                            ),

                            "Effect": "Undefined",

                            "Better Method": (
                                "Undefined"
                            ),
                        }
                    )

                    continue

                # =================================================
                # Paired values
                # =================================================

                a = (
                    conditional[
                        "A"
                    ]
                    .to_numpy(
                        dtype=float
                    )
                )

                b = (
                    conditional[
                        "B"
                    ]
                    .to_numpy(
                        dtype=float
                    )
                )

                differences = (
                    a
                    -
                    b
                )

                ties = np.isclose(
                    differences,
                    0.0,
                    atol=tolerance,
                    rtol=0.0,
                )

                # =================================================
                # Signed-rank tests
                # =================================================

                (
                    wilcox_stat,
                    wilcox_raw_p,
                ) = _run_wilcoxon(
                    a,
                    b,
                    zero_tolerance=(
                        tolerance
                    ),
                )

                (
                    pratt_stat,
                    pratt_raw_p,
                ) = _run_pratt(
                    a,
                    b,
                    zero_tolerance=(
                        tolerance
                    ),
                )

                (
                    zsplit_stat,
                    zsplit_raw_p,
                ) = _run_zsplit(
                    a,
                    b,
                    zero_tolerance=(
                        tolerance
                    ),
                )

                # =================================================
                # Effect size
                # =================================================

                effect = (
                    _rank_biserial_correlation(
                        a,
                        b,
                        zero_tolerance=(
                            tolerance
                        ),
                    )
                )

                better = (
                    _better_method_from_effect(
                        method_a,
                        method_b,
                        effect,
                        LOWER_IS_BETTER[
                            metric_name
                        ],
                        tolerance=tolerance,
                    )
                )

                # =================================================
                # Store
                # =================================================

                rows.append(
                    {
                        **base_row,

                        "N Retained": int(
                            len(
                                conditional
                            )
                        ),

                        "Retained %": (
                            len(
                                conditional
                            )
                            /
                            counts[
                                "N Paired"
                            ]
                            if counts[
                                "N Paired"
                            ]
                            else np.nan
                        ),

                        "N Non-Ties": int(
                            (
                                ~ties
                            ).sum()
                        ),

                        "Tie %": float(
                            ties.mean()
                        ),

                        "A Mean ± SD": (
                            _mean_sd_text(
                                a
                            )
                        ),

                        "B Mean ± SD": (
                            _mean_sd_text(
                                b
                            )
                        ),

                        "A Median [Q1, Q3]": (
                            _median_iqr_text(
                                a
                            )
                        ),

                        "B Median [Q1, Q3]": (
                            _median_iqr_text(
                                b
                            )
                        ),

                        "Median Difference A-B": (
                            float(
                                np.median(
                                    differences
                                )
                            )
                        ),

                        "Wilcox Statistic": (
                            wilcox_stat
                        ),

                        "Wilcox Raw p": (
                            wilcox_raw_p
                        ),

                        "Pratt Statistic": (
                            pratt_stat
                        ),

                        "Pratt Raw p": (
                            pratt_raw_p
                        ),

                        "Z-split Statistic": (
                            zsplit_stat
                        ),

                        "Z-split Raw p": (
                            zsplit_raw_p
                        ),

                        "Rank-Biserial (Non-Ties)": (
                            effect
                        ),

                        "Effect": (
                            _effect_magnitude(
                                effect
                            )
                        ),

                        "Better Method": (
                            better
                        ),
                    }
                )

    # ========================================================
    # DataFrame
    # ========================================================

    results = pd.DataFrame(
        rows
    )

    # ========================================================
    # Holm-corrected p-values
    # ========================================================

    results[
        "Wilcox Holm-adjusted p"
    ] = np.nan

    results[
        "Pratt Holm-adjusted p"
    ] = np.nan

    results[
        "Z-split Holm-adjusted p"
    ] = np.nan

    # --------------------------------------------------------
    # IMPORTANT:
    #
    # These are three sensitivity analyses of the SAME
    # hypotheses. We correct each zero-handling convention
    # independently, rather than treating the three p-values
    # as three separate substantive hypotheses.
    # --------------------------------------------------------

    for (
        _,
        group,
    ) in results.groupby(
        [
            "Condition",
            "Metric",
        ],
        sort=False,
    ):

        for (
            raw_col,
            adjusted_col,
        ) in [
            (
                "Wilcox Raw p",
                "Wilcox Holm-adjusted p",
            ),
            (
                "Pratt Raw p",
                "Pratt Holm-adjusted p",
            ),
            (
                "Z-split Raw p",
                "Z-split Holm-adjusted p",
            ),
        ]:

            valid = group.index[
                np.isfinite(
                    results.loc[
                        group.index,
                        raw_col,
                    ]
                )
            ]

            if len(
                valid
            ) == 0:

                continue

            (
                _,
                adjusted,
                _,
                _,
            ) = multipletests(
                results.loc[
                    valid,
                    raw_col,
                ].to_numpy(
                    dtype=float
                ),
                alpha=alpha,
                method="holm",
            )

            results.loc[
                valid,
                adjusted_col,
            ] = adjusted

    # ========================================================
    # Formatted Holm p-values
    # ========================================================

    results[
        "Wilcox Holm p"
    ] = (
        results[
            "Wilcox Holm-adjusted p"
        ]
        .apply(
            _format_p_value
        )
    )

    results[
        "Pratt Holm p"
    ] = (
        results[
            "Pratt Holm-adjusted p"
        ]
        .apply(
            _format_p_value
        )
    )

    results[
        "Z-split Holm p"
    ] = (
        results[
            "Z-split Holm-adjusted p"
        ]
        .apply(
            _format_p_value
        )
    )

    # ========================================================
    # Conclusions / sensitivity
    # ========================================================

    conclusions = []
    robustness = []

    for _, row in (
        results.iterrows()
    ):

        wilcox_p = row[
            "Wilcox Holm-adjusted p"
        ]

        pratt_p = row[
            "Pratt Holm-adjusted p"
        ]

        zsplit_p = row[
            "Z-split Holm-adjusted p"
        ]

        effect = row[
            "Rank-Biserial (Non-Ties)"
        ]

        better = row[
            "Better Method"
        ]

        # ----------------------------------------------------
        # Significance under each zero convention
        # ----------------------------------------------------

        wilcox_sig = (
            np.isfinite(
                wilcox_p
            )
            and
            wilcox_p < alpha
        )

        pratt_sig = (
            np.isfinite(
                pratt_p
            )
            and
            pratt_p < alpha
        )

        zsplit_sig = (
            np.isfinite(
                zsplit_p
            )
            and
            zsplit_p < alpha
        )

        # ----------------------------------------------------
        # Practical / effect threshold
        # ----------------------------------------------------

        practical = (
            np.isfinite(
                effect
            )
            and
            abs(
                effect
            )
            >=
            practical_effect_threshold
        )

        # ====================================================
        # Zero-handling robustness
        # ====================================================

        significance_results = [
            wilcox_sig,
            pratt_sig,
            zsplit_sig,
        ]

        if all(
            significance_results
        ):

            robustness.append(
                "Robust: all significant"
            )

        elif not any(
            significance_results
        ):

            robustness.append(
                "Robust: none significant"
            )

        else:

            robustness.append(
                "Sensitive to zero handling"
            )

        # ====================================================
        # Primary conclusion
        #
        # Keep your existing behaviour:
        # Wilcox remains the primary analysis.
        #
        # Pratt and Z-split are reported as sensitivity
        # analyses.
        # ====================================================

        if (
            wilcox_sig
            and
            practical
            and
            better
            not in {
                "Tie",
                "Undefined",
            }
        ):

            conclusion = (
                f"{better} better"
            )

        elif (
            wilcox_sig
            and
            not practical
        ):

            conclusion = (
                "Statistically significant; "
                "negligible effect"
            )

        elif (
            not wilcox_sig
            and
            practical
        ):

            conclusion = (
                "Non-negligible effect; "
                "not statistically significant"
            )

        else:

            conclusion = (
                "No meaningful evidence "
                "of difference"
            )

        conclusions.append(
            conclusion
        )

    results[
        "Zero-Handling Robustness"
    ] = robustness

    results[
        "Conclusion"
    ] = conclusions

    # ========================================================
    # Final output columns
    # ========================================================

    return results[
        [
            "Condition",
            "Metric",
            "Comparison",

            "N Paired Before Condition",
            "N Retained",
            "Retained %",

            "N Non-Ties",
            "Tie %",

            "A Mean ± SD",
            "B Mean ± SD",

            "A Median [Q1, Q3]",
            "B Median [Q1, Q3]",

            "Median Difference A-B",

            "Wilcox Holm p",
            "Pratt Holm p",
            "Z-split Holm p",

            "Rank-Biserial (Non-Ties)",
            "Effect",

            "Better Method",

            "Zero-Handling Robustness",
            "Conclusion",
        ]
    ].copy()


# ============================================================
# Trace length column detection
# ============================================================

def _resolve_trace_length_column(
    df: pd.DataFrame,
    requested: Optional[str] = None,
) -> str:

    if requested is not None:

        if requested not in df.columns:
            raise KeyError(
                f"Trace-length column "
                f"'{requested}' was not found."
            )

        return requested

    candidates = [
        "trace_length",
        "Trace Length",
        "trace_len",
        "prefix_length",
        "Prefix Length",
        "length",
    ]

    for candidate in candidates:

        if candidate in df.columns:
            return candidate

    raise KeyError(
        "Could not determine the trace-length column. "
        "Pass trace_length_column explicitly."
    )


# ============================================================
# Aggregate trace length
# ============================================================

def _aggregate_trace_length(
    combined: pd.DataFrame,
    *,
    trace_length_column: Optional[str] = None,
) -> pd.DataFrame:

    column = _resolve_trace_length_column(
        combined,
        trace_length_column,
    )

    df = combined[
        [
            *ID_COLUMNS,
            column,
        ]
    ].copy()

    df[
        column
    ] = pd.to_numeric(
        df[column],
        errors="coerce",
    )

    df = df.loc[
        np.isfinite(
            df[
                column
            ]
        )
    ]

    result = (
        df.groupby(
            ID_COLUMNS,
            dropna=False,
            as_index=False,
        )
        .agg(
            Trace_Length=(
                column,
                "median",
            )
        )
    )

    result[
        "Trace_Length"
    ] = (
        result[
            "Trace_Length"
        ]
        .round()
        .astype(int)
    )

    return result


# ============================================================
# Propagation-penalty outcome by trace length
# ============================================================

def build_propagation_penalty_by_trace_length(
    all_dataset_results: Mapping,
    *,
    baseline_method: str = "DiCE4EL+",
    condition: str = "Flip Success",
    tolerance: float = 1e-8,
    trace_length_column: Optional[str] = None,
) -> pd.DataFrame:

    if baseline_method not in {
        "DiCE4EL",
        "DiCE4EL+",
    }:
        raise ValueError(
            "baseline_method must be "
            "'DiCE4EL' or 'DiCE4EL+'."
        )

    rows = []

    for (
        dataset,
        seed_results,
    ) in all_dataset_results.items():

        combined = (
            _combine_method_results_across_seeds(
                seed_results
            )
        )

        retained, _ = (
            _get_common_success_queries(
                combined,
                condition,
                "ProACTS",
                baseline_method,
                tolerance,
            )
        )

        paired_penalty = _pair_metric(
            combined,
            "ProACTS",
            baseline_method,
            "process_violation",
            "process_violation",
        )

        trace_lengths = (
            _aggregate_trace_length(
                combined,
                trace_length_column=(
                    trace_length_column
                ),
            )
        )

        paired = (
            retained
            .merge(
                paired_penalty,
                on=ID_COLUMNS,
                how="inner",
                validate="one_to_one",
            )
            .merge(
                trace_lengths,
                on=ID_COLUMNS,
                how="inner",
                validate="one_to_one",
            )
        )

        classes = (
            _classify_propagation_penalty_pairs(
                paired[
                    "A"
                ].to_numpy(
                    dtype=float
                ),
                paired[
                    "B"
                ].to_numpy(
                    dtype=float
                ),
                tolerance=tolerance,
            )
        )

        paired = pd.concat(
            [
                paired.reset_index(
                    drop=True
                ),
                classes,
            ],
            axis=1,
        )

        for (
            trace_length,
            group,
        ) in paired.groupby(
            "Trace_Length",
            sort=True,
        ):

            n = len(
                group
            )

            for outcome in [
                "ProACTS better",
                "Tie",
                "Baseline better",
            ]:

                count = int(
                    (
                        group[
                            "Outcome"
                        ]
                        == outcome
                    ).sum()
                )

                rows.append(
                    {
                        "Dataset": dataset,
                        "Baseline": baseline_method,
                        "Condition": condition,
                        "Trace Length": int(
                            trace_length
                        ),
                        "Outcome": outcome,
                        "Count": count,
                        "N": n,
                        "Percent": (
                            count / n
                            if n
                            else np.nan
                        ),
                    }
                )

    return pd.DataFrame(
        rows
    )


# ============================================================
# Paired propagation-penalty plot by dataset
# ============================================================

def plot_paired_propagation_penalty_by_dataset(
    all_dataset_results,
    *,
    baseline_method="DiCE4EL+",
    condition="Flip Success",
    tolerance=1e-8,
    output_path=(
        "summary_evaluation/"
        "paired_propagation_penalty.pdf"
    ),
    show=True,
):
    """
    Paired propagation-penalty scatter plot by dataset.

    Each point:
        one factual query after taking the median
        result across seeds.

    Only queries where BOTH ProACTS and the baseline satisfy
    the specified success condition are included.

    x-axis:
        baseline propagation penalty

    y-axis:
        ProACTS propagation penalty

    Interpretation:
        below y=x -> ProACTS better
        on y=x    -> tie
        above y=x -> baseline better
    """

    # ========================================================
    # Validation
    # ========================================================

    if baseline_method not in {
        "DiCE4EL",
        "DiCE4EL+",
    }:
        raise ValueError(
            "baseline_method must be either "
            "'DiCE4EL' or 'DiCE4EL+'."
        )

    datasets = list(
        all_dataset_results.keys()
    )

    if not datasets:
        raise ValueError(
            "No dataset results supplied."
        )

    # ========================================================
    # Layout
    # ========================================================

    n_datasets = len(
        datasets
    )

    n_cols = min(
        5,
        n_datasets,
    )

    n_rows = math.ceil(
        n_datasets
        /
        n_cols
    )

    fig, axes = plt.subplots(
        n_rows,
        n_cols,
        figsize=(
            3.4 * n_cols,
            3.2 * n_rows,
        ),
        squeeze=False,
    )

    axes = axes.flatten()

    summary_rows = []

    # ========================================================
    # Colors
    # ========================================================

    # Okabe-Ito blue
    color_proacts = "#0072B2"

    # Neutral grey
    color_tie = "#8A8A8A"

    # Okabe-Ito orange
    color_baseline = "#E69F00"

    # ========================================================
    # Plot each dataset
    # ========================================================

    for panel_idx, dataset in enumerate(
        datasets
    ):

        ax = axes[
            panel_idx
        ]

        seed_results = (
            all_dataset_results[
                dataset
            ]
        )

        combined = (
            _combine_method_results_across_seeds(
                seed_results
            )
        )

        # ====================================================
        # Keep only queries where BOTH methods succeeded
        # ========================================================

        retained_ids, counts = (
            _get_common_success_queries(
                combined=combined,
                condition=condition,
                method_a="ProACTS",
                method_b=baseline_method,
                tolerance=tolerance,
            )
        )

        # ====================================================
        # Pair propagation penalties
        #
        # A = ProACTS
        # B = baseline
        # ========================================================

        paired_penalty = _pair_metric(
            combined=combined,
            method_a="ProACTS",
            method_b=baseline_method,
            column_a="process_violation",
            column_b="process_violation",
        )

        paired_penalty = (
            retained_ids.merge(
                paired_penalty,
                on=ID_COLUMNS,
                how="inner",
                validate="one_to_one",
            )
        )

        # ====================================================
        # No common successful queries
        # ========================================================

        if paired_penalty.empty:

            ax.text(
                0.5,
                0.5,
                "No common\nsuccessful queries",
                horizontalalignment="center",
                verticalalignment="center",
                transform=ax.transAxes,
                fontsize=10,
            )

            ax.set_title(
                dataset,
                fontsize=12,
                pad=5,
            )

            ax.set_axis_off()

            continue

        # ====================================================
        # Coordinates
        #
        # x = baseline
        # y = ProACTS
        # ========================================================

        x = (
            paired_penalty[
                "B"
            ]
            .to_numpy(
                dtype=float
            )
        )

        y = (
            paired_penalty[
                "A"
            ]
            .to_numpy(
                dtype=float
            )
        )

        # ====================================================
        # Classification
        # ========================================================

        proacts_better = (
            y
            <
            x - tolerance
        )

        baseline_better = (
            y
            >
            x + tolerance
        )

        ties = ~(
            proacts_better
            |
            baseline_better
        )

        # ====================================================
        # Scatter
        # ========================================================

        # Ties
        ax.scatter(
            x[
                ties
            ],
            y[
                ties
            ],
            alpha=0.32,
            s=25,
            color=color_tie,
            label="Tie",
            zorder=1,
        )

        # ProACTS better
        ax.scatter(
            x[
                proacts_better
            ],
            y[
                proacts_better
            ],
            alpha=0.55,
            s=26,
            color=color_proacts,
            label="ProACTS better",
            zorder=2,
        )

        # Baseline better
        ax.scatter(
            x[
                baseline_better
            ],
            y[
                baseline_better
            ],
            alpha=0.55,
            s=26,
            color=color_baseline,
            label=(
                f"{baseline_method} better"
            ),
            zorder=2,
        )

        # ====================================================
        # Common x/y scale
        # ========================================================

        combined_values = np.concatenate(
            [
                x,
                y,
            ]
        )

        value_min = float(
            np.min(
                combined_values
            )
        )

        value_max = float(
            np.max(
                combined_values
            )
        )

        lower = min(
            0.0,
            value_min,
        )

        upper = value_max

        if np.isclose(
            lower,
            upper,
        ):
            upper = (
                lower
                +
                1.0
            )

        # ----------------------------------------------------
        # Small padding around data
        # ----------------------------------------------------

        padding = (
            0.04
            *
            (
                upper
                -
                lower
            )
        )

        lower_plot = (
            lower
            -
            padding
        )

        upper_plot = (
            upper
            +
            padding
        )

        ax.set_xlim(
            lower_plot,
            upper_plot,
        )

        ax.set_ylim(
            lower_plot,
            upper_plot,
        )

        # ====================================================
        # Equality line
        # ========================================================

        ax.plot(
            [
                lower_plot,
                upper_plot,
            ],
            [
                lower_plot,
                upper_plot,
            ],
            linestyle="--",
            linewidth=1.1,
            color="0.35",
            zorder=0,
        )

        # Same physical scale on x and y
        ax.set_aspect(
            "equal",
            adjustable="box",
        )

        # ====================================================
        # Counts
        # ========================================================

        n = int(
            len(
                x
            )
        )

        n_proacts_better = int(
            proacts_better.sum()
        )

        n_ties = int(
            ties.sum()
        )

        n_baseline_better = int(
            baseline_better.sum()
        )

        # ====================================================
        # Dataset title = 12 pt
        # ========================================================

        ax.set_title(
            f"{dataset}\n"
            f"N={n}",
            fontsize=12,
            pad=5,
        )

        # ====================================================
        # Win / tie / loss annotation = 8.5 pt
        # ========================================================

        annotation = (
            f"ProACTS lower: "
            f"{100 * n_proacts_better / n:.1f}%\n"

            f"Tie: "
            f"{100 * n_ties / n:.1f}%\n"

            f"{baseline_method} lower: "
            f"{100 * n_baseline_better / n:.1f}%"
        )

        ax.text(
            0.04,
            0.96,
            annotation,
            transform=ax.transAxes,
            verticalalignment="top",
            fontsize=8.5,
        )

        # ====================================================
        # Grid / ticks
        # ========================================================

        ax.grid(
            alpha=0.18,
            linewidth=0.6,
        )

        ax.set_axisbelow(
            True
        )

        # Tick labels / numbers = 9 pt
        ax.tick_params(
            axis="both",
            labelsize=9,
            length=4,
            width=0.8,
        )

        # ====================================================
        # Numerical summary
        # ========================================================

        summary_rows.append(
            {
                "Dataset": dataset,

                "Condition": condition,

                "Baseline": baseline_method,

                "N Paired Before Condition": (
                    counts[
                        "N Paired"
                    ]
                ),

                "N Retained": n,

                "ProACTS Lower": (
                    n_proacts_better
                ),

                "Tie": (
                    n_ties
                ),

                f"{baseline_method} Lower": (
                    n_baseline_better
                ),

                "ProACTS Lower %": (
                    n_proacts_better
                    /
                    n
                ),

                "Tie %": (
                    n_ties
                    /
                    n
                ),

                f"{baseline_method} Lower %": (
                    n_baseline_better
                    /
                    n
                ),

                (
                    "Median ProACTS "
                    "Propagation Penalty"
                ): float(
                    np.median(
                        y
                    )
                ),

                (
                    f"Median {baseline_method} "
                    "Propagation Penalty"
                ): float(
                    np.median(
                        x
                    )
                ),
            }
        )

    # ========================================================
    # Hide unused panels
    # ========================================================

    for idx in range(
        n_datasets,
        len(
            axes
        ),
    ):
        axes[
            idx
        ].set_visible(
            False
        )

    # ========================================================
    # Shared axis titles = 14 pt
    # ========================================================

    fig.supxlabel(
        f"{baseline_method} Propagation Loss",
        fontsize=14,
    )

    fig.supylabel(
        "ProACTS Propagation Loss",
        fontsize=14,
    )

    # ========================================================
    # Shared legend
    # ========================================================

    legend_handles = [

        plt.Line2D(
            [0],
            [0],
            marker="o",
            linestyle="",
            markersize=8,
            markerfacecolor=color_proacts,
            markeredgecolor=color_proacts,
            label="ProACTS better",
        ),

        plt.Line2D(
            [0],
            [0],
            marker="o",
            linestyle="",
            markersize=8,
            markerfacecolor=color_tie,
            markeredgecolor=color_tie,
            label="Tie",
        ),

        plt.Line2D(
            [0],
            [0],
            marker="o",
            linestyle="",
            markersize=8,
            markerfacecolor=color_baseline,
            markeredgecolor=color_baseline,
            label=(
                f"{baseline_method} better"
            ),
        ),
    ]

    # ========================================================
    # Legend = 12 pt
    # ========================================================

    fig.legend(
        handles=legend_handles,
        loc="upper center",
        ncol=3,
        frameon=False,
        bbox_to_anchor=(
            0.5,
            0.985,
        ),
        fontsize=12,
        columnspacing=1.8,
        handletextpad=0.6,
    )

    # ========================================================
    # Manual spacing
    #
    # Keep extra vertical space because dataset titles have
    # two lines and the legend occupies the top area.
    # ========================================================

    fig.subplots_adjust(
        left=0.065,
        right=0.985,
        bottom=0.10,
        top=0.84,
        wspace=0.12,
        hspace=0.32,
    )

    # ========================================================
    # Save
    # ========================================================

    if output_path is not None:

        directory = os.path.dirname(
            output_path
        )

        if directory:
            os.makedirs(
                directory,
                exist_ok=True,
            )

        fig.savefig(
            output_path,
            bbox_inches="tight",
        )

    # ========================================================
    # Show / close
    # ========================================================

    if show:
        plt.show()

    else:
        plt.close(
            fig
        )

    return pd.DataFrame(
        summary_rows
    )


def plot_propagation_penalty_outcome_by_trace_length(
    trace_length_results: pd.DataFrame,
    *,
    baseline_method: str = "DiCE4EL+",
    output_path: Optional[str] = (
        "summary_evaluation/"
        "propagation_penalty_outcome_by_trace_length.pdf"
    ),
    show: bool = True,
) -> None:

    datasets = (
        trace_length_results[
            "Dataset"
        ]
        .drop_duplicates()
        .tolist()
    )

    if not datasets:
        raise ValueError(
            "No trace-length results."
        )

    n_cols = 5

    n_rows = math.ceil(
        len(datasets)
        /
        n_cols
    )

    fig, axes = plt.subplots(
        n_rows,
        n_cols,
        figsize=(
            3.5 * n_cols,
            3.1 * n_rows,
        ),
        sharey=True,
        squeeze=False,
    )

    axes = axes.flatten()

    display_labels = {
        "ProACTS better": "ProACTS better",
        "Tie": "Tie",
        "Baseline better": (
            f"{baseline_method} better"
        ),
    }

    for idx, dataset in enumerate(
        datasets
    ):

        ax = axes[
            idx
        ]

        dataset_df = (
            trace_length_results.loc[
                trace_length_results[
                    "Dataset"
                ].eq(dataset)
            ]
            .copy()
        )

        for outcome in [
            "ProACTS better",
            "Tie",
            "Baseline better",
        ]:

            subset = (
                dataset_df.loc[
                    dataset_df[
                        "Outcome"
                    ].eq(outcome)
                ]
                .sort_values(
                    "Trace Length"
                )
            )

            if subset.empty:
                continue

            ax.plot(
                subset[
                    "Trace Length"
                ],
                subset[
                    "Percent"
                ],
                marker="o",
                markersize=4,
                linewidth=1.8,
                color=(
                    PROPAGATION_OUTCOME_COLORS[
                        outcome
                    ]
                ),
                label=(
                    display_labels[
                        outcome
                    ]
                ),
            )

        ax.set_title(
            dataset
        )

        ax.set_ylim(
            0,
            1,
        )

        ax.grid(
            alpha=0.20
        )

    for idx in range(
        len(datasets),
        len(axes),
    ):
        axes[
            idx
        ].set_visible(
            False
        )

    fig.supxlabel(
        "Trace length"
    )

    fig.supylabel(
        "Proportion of paired successful queries"
    )

    handles, labels = (
        axes[
            0
        ].get_legend_handles_labels()
    )

    fig.legend(
        handles,
        labels,
        loc="upper center",
        ncol=3,
        frameon=False,
        bbox_to_anchor=(
            0.5,
            1.01,
        ),
    )

    fig.tight_layout(
        rect=[
            0,
            0,
            1,
            0.95,
        ]
    )

    if output_path:

        directory = os.path.dirname(
            output_path
        )

        if directory:
            os.makedirs(
                directory,
                exist_ok=True,
            )

        fig.savefig(
            output_path,
            bbox_inches="tight",
        )

    if show:
        plt.show()

    else:
        plt.close(
            fig
        )


# ============================================================
# ProACTS multiple - aggregate per query
# ============================================================

def _aggregate_proacts_multiple_metric(
    combined: pd.DataFrame,
    metric_column: str,
) -> pd.DataFrame:

    required = [
        "Seed",
        *ID_COLUMNS,
        DESIRED_COUNT_COLUMN,
        metric_column,
    ]

    missing = [
        c
        for c in required
        if c not in combined.columns
    ]

    if missing:
        raise KeyError(
            f"Missing columns for "
            f"'{metric_column}': {missing}"
        )

    df = combined[
        required
    ].copy()

    df[
        metric_column
    ] = pd.to_numeric(
        df[
            metric_column
        ],
        errors="coerce",
    )

    df = df.loc[
        np.isfinite(
            df[
                metric_column
            ]
        )
    ]

    return (
        df.groupby(
            [
                *ID_COLUMNS,
                DESIRED_COUNT_COLUMN,
            ],
            dropna=False,
            as_index=False,
        )
        .agg(
            Value=(
                metric_column,
                "median",
            ),
            N_Seeds=(
                "Seed",
                "nunique",
            ),
        )
    )


# ============================================================
# ProACTS multiple main results
# ============================================================

def _build_proacts_multiple_main_results(
    combined: pd.DataFrame,
) -> pd.DataFrame:

    rows = []

    for (
        metric_name,
        source_column,
    ) in PROACTS_MULTIPLE_METRICS.items():

        aggregated = (
            _aggregate_proacts_multiple_metric(
                combined,
                source_column,
            )
        )

        for (
            desired_count,
            group,
        ) in aggregated.groupby(
            DESIRED_COUNT_COLUMN,
            sort=True,
        ):

            values = (
                group[
                    "Value"
                ]
                .to_numpy(
                    dtype=float
                )
            )

            rows.append(
                {
                    "Desired Positions": int(
                        desired_count
                    ),
                    "Metric": metric_name,
                    "N Instances": len(
                        values
                    ),
                    "Mean ± SD": (
                        _mean_sd_text(
                            values
                        )
                    ),
                    "Median [Q1, Q3]": (
                        _median_iqr_text(
                            values
                        )
                    ),
                }
            )

    return (
        pd.DataFrame(
            rows
        )
        .sort_values(
            [
                "Desired Positions",
                "Metric",
            ]
        )
        .reset_index(
            drop=True
        )
    )


# ============================================================
# ProACTS multiple flip rate
# ============================================================

def _build_proacts_multiple_flip_rate(
    combined: pd.DataFrame,
    *,
    tolerance: float = 1e-8,
    flip_column: str = "margin_loss_flipped",
) -> pd.DataFrame:

    df = combined[
        [
            "Seed",
            DESIRED_COUNT_COLUMN,
            flip_column,
        ]
    ].copy()

    df[
        flip_column
    ] = pd.to_numeric(
        df[
            flip_column
        ],
        errors="coerce",
    )

    df = df.loc[
        np.isfinite(
            df[
                flip_column
            ]
        )
    ]

    raw_rows = []

    for (
        desired_count,
        seed,
    ), group in df.groupby(
        [
            DESIRED_COUNT_COLUMN,
            "Seed",
        ],
        sort=True,
    ):

        values = (
            group[
                flip_column
            ]
            .to_numpy(
                dtype=float
            )
        )

        success = (
            values <= tolerance
        )

        raw_rows.append(
            {
                "Desired Positions": int(
                    desired_count
                ),
                "Seed": seed,
                "N Instances": len(
                    values
                ),
                "Successful": int(
                    success.sum()
                ),
                "Flip Rate": float(
                    success.mean()
                ),
            }
        )

    raw = pd.DataFrame(
        raw_rows
    )

    rows = []

    for (
        desired_count,
        group,
    ) in raw.groupby(
        "Desired Positions",
        sort=True,
    ):

        rates = (
            group[
                "Flip Rate"
            ]
            .to_numpy(
                dtype=float
            )
        )

        rows.append(
            {
                "Desired Positions": int(
                    desired_count
                ),
                "N Seeds": len(
                    rates
                ),
                "Mean Flip Rate": float(
                    np.mean(
                        rates
                    )
                ),
                "Mean ± SD": (
                    _mean_sd_text(
                        rates
                    )
                ),
            }
        )

    return pd.DataFrame(
        rows
    )


# ============================================================
# ProACTS multiple seed robustness
# ============================================================

def _build_proacts_multiple_seed_robustness(
    combined: pd.DataFrame,
) -> pd.DataFrame:

    rows = []

    for (
        metric_name,
        source_column,
    ) in PROACTS_MULTIPLE_METRICS.items():

        df = combined[
            [
                "Seed",
                DESIRED_COUNT_COLUMN,
                source_column,
            ]
        ].copy()

        df[
            source_column
        ] = pd.to_numeric(
            df[
                source_column
            ],
            errors="coerce",
        )

        df = df.loc[
            np.isfinite(
                df[
                    source_column
                ]
            )
        ]

        for (
            desired_count,
            desired_df,
        ) in df.groupby(
            DESIRED_COUNT_COLUMN,
            sort=True,
        ):

            seed_means = []
            seed_medians = []

            for _, seed_df in (
                desired_df.groupby(
                    "Seed",
                    sort=True,
                )
            ):

                values = (
                    seed_df[
                        source_column
                    ]
                    .to_numpy(
                        dtype=float
                    )
                )

                seed_means.append(
                    np.mean(
                        values
                    )
                )

                seed_medians.append(
                    np.median(
                        values
                    )
                )

            rows.append(
                {
                    "Desired Positions": int(
                        desired_count
                    ),
                    "Metric": metric_name,
                    "N Seeds": len(
                        seed_means
                    ),
                    "Seed Mean ± SD": (
                        _mean_sd_text(
                            np.asarray(
                                seed_means
                            )
                        )
                    ),
                    "Seed Median ± SD": (
                        _mean_sd_text(
                            np.asarray(
                                seed_medians
                            )
                        )
                    ),
                }
            )

    return pd.DataFrame(
        rows
    )


# ============================================================
# Excel styling
# ============================================================

def _autosize_workbook(
    output_path: str,
) -> None:

    from openpyxl import load_workbook

    from openpyxl.styles import (
        Alignment,
        Font,
        PatternFill,
    )

    workbook = load_workbook(
        output_path
    )

    header_fill = PatternFill(
        "solid",
        fgColor="D9EAF7",
    )

    for worksheet in workbook.worksheets:

        worksheet.freeze_panes = (
            "A2"
        )

        worksheet.auto_filter.ref = (
            worksheet.dimensions
        )

        for cell in worksheet[1]:

            cell.font = Font(
                bold=True
            )

            cell.fill = (
                header_fill
            )

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        for column_cells in (
            worksheet.columns
        ):

            values = [
                ""
                if cell.value is None
                else str(
                    cell.value
                )
                for cell
                in column_cells
            ]

            width = min(
                max(
                    len(
                        value
                    )
                    for value
                    in values
                )
                + 2,
                45,
            )

            worksheet.column_dimensions[
                column_cells[
                    0
                ].column_letter
            ].width = width

        for row in worksheet.iter_rows(
            min_row=2
        ):

            for cell in row:

                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

    workbook.save(
        output_path
    )


# ============================================================
# One dataset - single desired position
# ============================================================

def summarize_across_seeds(
    all_results: Mapping,
    *,
    eval_name: Optional[str] = None,
    output_dir: str = "summary_evaluation",
    alpha: float = 0.05,
    tolerance: float = 1e-8,
    practical_effect_threshold: float = 0.10,
) -> dict:

    combined = (
        _combine_method_results_across_seeds(
            all_results
        )
    )

    main_results = (
        _build_main_results_table(
            combined
        )
    )

    success_rates = (
        _build_success_rate_table(
            combined,
            tolerance,
        )
    )

    success_significance = (
        _build_success_significance_table(
            combined,
            alpha=alpha,
            tolerance=tolerance,
        )
    )

    seed_robustness = (
        _build_seed_robustness_table(
            combined
        )
    )

    retained_counts = (
        _build_retained_counts_table(
            combined,
            tolerance,
        )
    )

    conditional_results = (
        _build_conditional_pairwise_results(
            combined,
            alpha=alpha,
            tolerance=tolerance,
            practical_effect_threshold=(
                practical_effect_threshold
            ),
        )
    )

    propagation_ties = pd.concat(
        [
            _build_propagation_penalty_tie_table(
                combined,
                condition="Flip Success",
                baseline_method="DiCE4EL",
                tolerance=tolerance,
            ),

            _build_propagation_penalty_tie_table(
                combined,
                condition="Flip Success",
                baseline_method="DiCE4EL+",
                tolerance=tolerance,
            ),
        ],
        ignore_index=True,
    )

    results = {
        "main_results": main_results,
        "success_rates": success_rates,
        "success_significance": success_significance,
        "seed_robustness": seed_robustness,
        "retained_counts": retained_counts,
        "conditional_results": conditional_results,
        "propagation_penalty_ties": propagation_ties,
    }

    if eval_name is not None:

        os.makedirs(
            output_dir,
            exist_ok=True,
        )

        filename = (
            eval_name
            if eval_name.endswith(
                ".xlsx"
            )
            else f"{eval_name}.xlsx"
        )

        output_path = (
            os.path.join(
                output_dir,
                filename,
            )
        )

        with pd.ExcelWriter(
            output_path,
            engine="openpyxl",
        ) as writer:

            main_results.to_excel(
                writer,
                sheet_name="Main_Results",
                index=False,
            )

            success_rates.to_excel(
                writer,
                sheet_name="Success_Rates",
                index=False,
            )

            success_significance.to_excel(
                writer,
                sheet_name="Success_Significance",
                index=False,
            )

            seed_robustness.to_excel(
                writer,
                sheet_name="Seed_Robustness",
                index=False,
            )

            retained_counts.to_excel(
                writer,
                sheet_name="Retained_Counts",
                index=False,
            )

            conditional_results.to_excel(
                writer,
                sheet_name="Conditional_Results",
                index=False,
            )

            propagation_ties.to_excel(
                writer,
                sheet_name="Propagation_Penalty_Ties",
                index=False,
            )

        _autosize_workbook(
            output_path
        )

    return results


# ============================================================
# All datasets - single desired position
# ============================================================

def summarize_datasets(
    all_dataset_results: Mapping,
    *,
    output_path: str = (
        "summary_evaluation/"
        "all_datasets_results.xlsx"
    ),
    alpha: float = 0.05,
    tolerance: float = 1e-8,
    practical_effect_threshold: float = 0.10,
    trace_length_baseline: str = "DiCE4EL+",
    trace_length_condition: str = "Flip Success",
    trace_length_column: Optional[str] = None,
) -> dict:

    main_frames = []
    success_frames = []
    success_significance_frames = []
    robustness_frames = []
    retained_frames = []
    conditional_frames = []
    propagation_tie_frames = []

    for (
        dataset,
        seed_results,
    ) in all_dataset_results.items():

        tables = summarize_across_seeds(
            seed_results,
            eval_name=None,
            alpha=alpha,
            tolerance=tolerance,
            practical_effect_threshold=(
                practical_effect_threshold
            ),
        )

        mapping = {
            "main_results": main_frames,
            "success_rates": success_frames,
            "success_significance": (
                success_significance_frames
            ),
            "seed_robustness": robustness_frames,
            "retained_counts": retained_frames,
            "conditional_results": conditional_frames,
            "propagation_penalty_ties": (
                propagation_tie_frames
            ),
        }

        for key, target in mapping.items():

            frame = (
                tables[
                    key
                ].copy()
            )

            frame.insert(
                0,
                "Dataset",
                dataset,
            )

            target.append(
                frame
            )

    trace_length_results = (
        build_propagation_penalty_by_trace_length(
            all_dataset_results,
            baseline_method=(
                trace_length_baseline
            ),
            condition=(
                trace_length_condition
            ),
            tolerance=tolerance,
            trace_length_column=(
                trace_length_column
            ),
        )
    )

    outputs = {
        "main_results": pd.concat(
            main_frames,
            ignore_index=True,
        ),

        "success_rates": pd.concat(
            success_frames,
            ignore_index=True,
        ),

        "success_significance": pd.concat(
            success_significance_frames,
            ignore_index=True,
        ),

        "seed_robustness": pd.concat(
            robustness_frames,
            ignore_index=True,
        ),

        "retained_counts": pd.concat(
            retained_frames,
            ignore_index=True,
        ),

        "conditional_results": pd.concat(
            conditional_frames,
            ignore_index=True,
        ),

        "propagation_penalty_ties": pd.concat(
            propagation_tie_frames,
            ignore_index=True,
        ),

        "propagation_penalty_trace_length": (
            trace_length_results
        ),
    }

    directory = os.path.dirname(
        output_path
    )

    if directory:
        os.makedirs(
            directory,
            exist_ok=True,
        )

    with pd.ExcelWriter(
        output_path,
        engine="openpyxl",
    ) as writer:

        outputs[
            "main_results"
        ].to_excel(
            writer,
            sheet_name="Main_Results",
            index=False,
        )

        outputs[
            "success_rates"
        ].to_excel(
            writer,
            sheet_name="Success_Rates",
            index=False,
        )

        outputs[
            "success_significance"
        ].to_excel(
            writer,
            sheet_name="Success_Significance",
            index=False,
        )

        outputs[
            "seed_robustness"
        ].to_excel(
            writer,
            sheet_name="Seed_Robustness",
            index=False,
        )

        outputs[
            "retained_counts"
        ].to_excel(
            writer,
            sheet_name="Retained_Counts",
            index=False,
        )

        outputs[
            "conditional_results"
        ].to_excel(
            writer,
            sheet_name="Conditional_Results",
            index=False,
        )

        outputs[
            "propagation_penalty_ties"
        ].to_excel(
            writer,
            sheet_name="Propagation_Penalty_Ties",
            index=False,
        )

        outputs[
            "propagation_penalty_trace_length"
        ].to_excel(
            writer,
            sheet_name="Propagation_Trace_Length",
            index=False,
        )

    _autosize_workbook(
        output_path
    )

    return outputs


# ============================================================
# One dataset - multiple desired positions
# ============================================================

def summarize_proacts_multiple_positions(
    seed_results: Mapping,
    *,
    eval_name: Optional[str] = None,
    output_dir: str = "summary_evaluation",
    tolerance: float = 1e-8,
) -> dict:

    combined = (
        _combine_proacts_multiple_across_seeds(
            seed_results
        )
    )

    main_results = (
        _build_proacts_multiple_main_results(
            combined
        )
    )

    flip_rate = (
        _build_proacts_multiple_flip_rate(
            combined,
            tolerance=tolerance,
        )
    )

    seed_robustness = (
        _build_proacts_multiple_seed_robustness(
            combined
        )
    )

    results = {
        "main_results": main_results,
        "flip_rate": flip_rate,
        "seed_robustness": (
            seed_robustness
        ),
    }

    if eval_name is not None:

        os.makedirs(
            output_dir,
            exist_ok=True,
        )

        filename = (
            eval_name
            if eval_name.endswith(
                ".xlsx"
            )
            else f"{eval_name}.xlsx"
        )

        output_path = os.path.join(
            output_dir,
            filename,
        )

        with pd.ExcelWriter(
            output_path,
            engine="openpyxl",
        ) as writer:

            main_results.to_excel(
                writer,
                sheet_name="Main_Results",
                index=False,
            )

            flip_rate.to_excel(
                writer,
                sheet_name="Flip_Rate",
                index=False,
            )

            seed_robustness.to_excel(
                writer,
                sheet_name="Seed_Robustness",
                index=False,
            )

        _autosize_workbook(
            output_path
        )

    return results


# ============================================================
# All datasets - multiple desired positions
# ============================================================

def summarize_proacts_multiple_datasets(
    all_dataset_results: Mapping,
    *,
    output_path: str = (
        "summary_evaluation/"
        "proacts_multiple_positions.xlsx"
    ),
    tolerance: float = 1e-8,
) -> dict:

    main_frames = []
    flip_frames = []
    robustness_frames = []

    for (
        dataset,
        seed_results,
    ) in all_dataset_results.items():

        result = (
            summarize_proacts_multiple_positions(
                seed_results,
                eval_name=None,
                tolerance=tolerance,
            )
        )

        for key, target in [
            (
                "main_results",
                main_frames,
            ),
            (
                "flip_rate",
                flip_frames,
            ),
            (
                "seed_robustness",
                robustness_frames,
            ),
        ]:

            frame = (
                result[
                    key
                ].copy()
            )

            frame.insert(
                0,
                "Dataset",
                dataset,
            )

            target.append(
                frame
            )

    outputs = {
        "main_results": pd.concat(
            main_frames,
            ignore_index=True,
        ),

        "flip_rate": pd.concat(
            flip_frames,
            ignore_index=True,
        ),

        "seed_robustness": pd.concat(
            robustness_frames,
            ignore_index=True,
        ),
    }

    directory = os.path.dirname(
        output_path
    )

    if directory:
        os.makedirs(
            directory,
            exist_ok=True,
        )

    with pd.ExcelWriter(
        output_path,
        engine="openpyxl",
    ) as writer:

        outputs[
            "main_results"
        ].to_excel(
            writer,
            sheet_name="Main_Results",
            index=False,
        )

        outputs[
            "flip_rate"
        ].to_excel(
            writer,
            sheet_name="Flip_Rate",
            index=False,
        )

        outputs[
            "seed_robustness"
        ].to_excel(
            writer,
            sheet_name="Seed_Robustness",
            index=False,
        )

    _autosize_workbook(
        output_path
    )

    return outputs


# ============================================================
# Loss vs propagation penalty plots
# ============================================================

def _build_loss_propagation_penalty_method_frame(
    combined: pd.DataFrame,
    *,
    method: str,
    loss_column: str,
) -> pd.DataFrame:
    """
    Build one per-query loss/propagation-penalty frame
    for one method.

    The metric value for each query is first aggregated
    across seeds using the median.
    """

    loss = (
        _aggregate_method_metric_across_seeds(
            combined=combined,
            method=method,
            source_column=loss_column,
        )
        .rename(
            columns={
                "Value": "Loss"
            }
        )
        [[*ID_COLUMNS, "Loss"]]
    )

    penalty = (
        _aggregate_method_metric_across_seeds(
            combined=combined,
            method=method,
            source_column="process_violation",
        )
        .rename(
            columns={
                "Value": "PropagationPenalty"
            }
        )
        [[
            *ID_COLUMNS,
            "PropagationPenalty",
        ]]
    )

    return loss.merge(
        penalty,
        on=ID_COLUMNS,
        how="inner",
        validate="one_to_one",
    )


def plot_loss_vs_propagation_penalty_by_dataset(
    all_dataset_results,
    *,
    baseline_method="DiCE4EL+",
    loss_metric="Flip Loss",
    output_path=(
        "summary_evaluation/"
        "loss_vs_propagation_penalty.pdf"
    ),
    show=True,
):
    """
    Combined ProACTS-vs-baseline scatter plot.

    Each point is one factual query after taking the median
    across seeds.

    x-axis:
        Flip Loss or Margin Loss

    y-axis:
        Propagation Penalty

    ProACTS uses filled blue circles.
    The baseline uses hollow orange circles so overlapping
    points from the two techniques remain visible.
    """

    if baseline_method not in {
        "DiCE4EL",
        "DiCE4EL+",
    }:
        raise ValueError(
            "baseline_method must be either "
            "'DiCE4EL' or 'DiCE4EL+'."
        )

    if loss_metric == "Flip Loss":
        loss_columns = {
            "ProACTS": "margin_loss_flipped",
            "DiCE4EL": "margin_loss_na_flipped",
            "DiCE4EL+": "margin_loss_na_flipped",
        }
        x_label = "Flip loss"

    elif loss_metric == "Margin Loss":
        loss_columns = {
            "ProACTS": "margin_loss",
            "DiCE4EL": "margin_loss_na",
            "DiCE4EL+": "margin_loss_na",
        }
        x_label = "Margin loss"

    else:
        raise ValueError(
            "loss_metric must be either "
            "'Flip Loss' or 'Margin Loss'."
        )

    datasets = list(all_dataset_results.keys())

    if not datasets:
        raise ValueError("No dataset results supplied.")

    n_datasets = len(datasets)
    n_cols = min(5, n_datasets)
    n_rows = math.ceil(n_datasets / n_cols)

    fig, axes = plt.subplots(
        n_rows,
        n_cols,
        figsize=(3.4 * n_cols, 3.0 * n_rows),
        squeeze=False,
        sharey=False,
    )
    axes = axes.flatten()

    color_proacts = "#1f77b4"
    color_baseline = "#ff7f0e"

    summary_rows = []

    for panel_idx, dataset in enumerate(datasets):

        ax = axes[panel_idx]

        combined = _combine_method_results_across_seeds(
            all_dataset_results[dataset]
        )

        proacts = _build_loss_propagation_penalty_method_frame(
            combined,
            method="ProACTS",
            loss_column=loss_columns["ProACTS"],
        )

        baseline = _build_loss_propagation_penalty_method_frame(
            combined,
            method=baseline_method,
            loss_column=loss_columns[baseline_method],
        )

        proacts_x = proacts["Loss"].to_numpy(dtype=float)
        proacts_y = proacts["PropagationPenalty"].to_numpy(dtype=float)

        baseline_x = baseline["Loss"].to_numpy(dtype=float)
        baseline_y = baseline["PropagationPenalty"].to_numpy(dtype=float)

        # Filled ProACTS points.
        ax.scatter(
            proacts_x,
            proacts_y,
            s=22,
            alpha=0.42,
            color=color_proacts,
            marker="o",
            label="ProACTS",
            zorder=2,
        )

        # Hollow baseline points: overlap remains visible.
        ax.scatter(
            baseline_x,
            baseline_y,
            s=26,
            alpha=0.72,
            facecolors="none",
            edgecolors=color_baseline,
            linewidths=0.9,
            marker="o",
            label=baseline_method,
            zorder=3,
        )

        ax.set_title(dataset)
        ax.grid(alpha=0.20)

        all_penalty = np.concatenate([proacts_y, baseline_y])
        if len(all_penalty):
            penalty_max = float(np.max(all_penalty))
            ax.set_ylim(
                bottom=0,
                top=(penalty_max * 1.05 if penalty_max > 0 else 1.0),
            )

        summary_rows.append(
            {
                "Dataset": dataset,
                "Loss": loss_metric,
                "Baseline": baseline_method,
                "ProACTS N": int(len(proacts)),
                f"{baseline_method} N": int(len(baseline)),
                "ProACTS Median Loss": float(np.median(proacts_x)),
                "ProACTS Median Propagation Penalty": float(np.median(proacts_y)),
                f"{baseline_method} Median Loss": float(
                    np.median(baseline_x)
                ),
                f"{baseline_method} Median Propagation Penalty": float(
                    np.median(baseline_y)
                ),
            }
        )

    for idx in range(n_datasets, len(axes)):
        axes[idx].set_visible(False)

    fig.supxlabel(x_label)
    fig.supylabel("Propagation penalty")

    legend_handles = [
        plt.Line2D(
            [0],
            [0],
            marker="o",
            linestyle="",
            markersize=6,
            markerfacecolor=color_proacts,
            markeredgecolor=color_proacts,
            label="ProACTS",
        ),
        plt.Line2D(
            [0],
            [0],
            marker="o",
            linestyle="",
            markersize=6,
            markerfacecolor="none",
            markeredgecolor=color_baseline,
            label=baseline_method,
        ),
    ]

    fig.legend(
        handles=legend_handles,
        loc="upper center",
        ncol=2,
        frameon=False,
        bbox_to_anchor=(0.5, 1.01),
    )

    fig.tight_layout(rect=[0, 0, 1, 0.95])

    if output_path is not None:
        directory = os.path.dirname(output_path)
        if directory:
            os.makedirs(directory, exist_ok=True)
        fig.savefig(output_path, bbox_inches="tight")

    if show:
        plt.show()
    else:
        plt.close(fig)

    return pd.DataFrame(summary_rows)


def plot_loss_vs_propagation_penalty_single_method(
    all_dataset_results,
    *,
    method="ProACTS",
    loss_metric="Flip Loss",
    output_path=None,
    show=True,
):
    """
    Plot loss versus propagation penalty for ONE technique.

    This removes all between-method overlap and is useful for
    inspecting the within-technique loss/propagation-penalty relationship.
    """

    if method not in {
        "ProACTS",
        "DiCE4EL",
        "DiCE4EL+",
    }:
        raise ValueError(
            "method must be 'ProACTS', 'DiCE4EL', or 'DiCE4EL+'."
        )

    if loss_metric == "Flip Loss":
        loss_columns = {
            "ProACTS": "margin_loss_flipped",
            "DiCE4EL": "margin_loss_na_flipped",
            "DiCE4EL+": "margin_loss_na_flipped",
        }
        x_label = "Flip loss"

    elif loss_metric == "Margin Loss":
        loss_columns = {
            "ProACTS": "margin_loss",
            "DiCE4EL": "margin_loss_na",
            "DiCE4EL+": "margin_loss_na",
        }
        x_label = "Margin loss"

    else:
        raise ValueError(
            "loss_metric must be either "
            "'Flip Loss' or 'Margin Loss'."
        )

    datasets = list(all_dataset_results.keys())

    if not datasets:
        raise ValueError("No dataset results supplied.")

    n_datasets = len(datasets)
    n_cols = min(5, n_datasets)
    n_rows = math.ceil(n_datasets / n_cols)

    fig, axes = plt.subplots(
        n_rows,
        n_cols,
        figsize=(3.4 * n_cols, 3.0 * n_rows),
        squeeze=False,
        sharey=False,
    )
    axes = axes.flatten()

    method_colors = {
        "ProACTS": "#1f77b4",
        "DiCE4EL": "#ff7f0e",
        "DiCE4EL+": "#ff7f0e",
    }
    color = method_colors[method]

    summary_rows = []

    for panel_idx, dataset in enumerate(datasets):

        ax = axes[panel_idx]

        combined = _combine_method_results_across_seeds(
            all_dataset_results[dataset]
        )

        frame = _build_loss_propagation_penalty_method_frame(
            combined,
            method=method,
            loss_column=loss_columns[method],
        )

        x = frame["Loss"].to_numpy(dtype=float)
        y = frame["PropagationPenalty"].to_numpy(dtype=float)

        ax.scatter(
            x,
            y,
            s=20,
            alpha=0.42,
            color=color,
        )

        ax.set_title(dataset)
        ax.grid(alpha=0.20)

        if len(y):
            penalty_max = float(np.max(y))
            ax.set_ylim(
                bottom=0,
                top=(penalty_max * 1.05 if penalty_max > 0 else 1.0),
            )

        summary_rows.append(
            {
                "Dataset": dataset,
                "Method": method,
                "Loss": loss_metric,
                "N": int(len(frame)),
                "Median Loss": float(np.median(x)),
                "Median Propagation Penalty": float(np.median(y)),
            }
        )

    for idx in range(n_datasets, len(axes)):
        axes[idx].set_visible(False)

    fig.supxlabel(x_label)
    fig.supylabel("Propagation penalty")

    fig.legend(
        handles=[
            plt.Line2D(
                [0],
                [0],
                marker="o",
                linestyle="",
                markersize=6,
                color=color,
                label=method,
            )
        ],
        loc="upper center",
        ncol=1,
        frameon=False,
        bbox_to_anchor=(0.5, 1.01),
    )

    fig.tight_layout(rect=[0, 0, 1, 0.95])

    if output_path is not None:
        directory = os.path.dirname(output_path)
        if directory:
            os.makedirs(directory, exist_ok=True)
        fig.savefig(output_path, bbox_inches="tight")

    if show:
        plt.show()
    else:
        plt.close(fig)

    return pd.DataFrame(summary_rows)


def generate_loss_propagation_penalty_plots(
    all_dataset_results,
    *,
    baseline_method="DiCE4EL+",
    output_dir="summary_evaluation",
    show=True,
):
    """
    Generate BOTH:
      1. combined ProACTS vs baseline plots, and
      2. separate ProACTS and baseline plots,

    for BOTH Flip Loss and Margin Loss.
    """

    if baseline_method not in {
        "DiCE4EL",
        "DiCE4EL+",
    }:
        raise ValueError(
            "baseline_method must be either "
            "'DiCE4EL' or 'DiCE4EL+'."
        )

    os.makedirs(output_dir, exist_ok=True)

    safe_baseline = (
        baseline_method
        .replace("+", "plus")
        .replace(" ", "_")
        .lower()
    )

    outputs = {}

    for loss_metric, prefix in [
        ("Flip Loss", "flip_loss"),
        ("Margin Loss", "margin_loss"),
    ]:

        outputs[f"{prefix}_combined"] = (
            plot_loss_vs_propagation_penalty_by_dataset(
                all_dataset_results,
                baseline_method=baseline_method,
                loss_metric=loss_metric,
                output_path=os.path.join(
                    output_dir,
                    f"{prefix}_vs_propagation_penalty_"
                    f"proacts_vs_{safe_baseline}.pdf",
                ),
                show=show,
            )
        )

        outputs[f"{prefix}_proacts"] = (
            plot_loss_vs_propagation_penalty_single_method(
                all_dataset_results,
                method="ProACTS",
                loss_metric=loss_metric,
                output_path=os.path.join(
                    output_dir,
                    f"{prefix}_vs_propagation_penalty_proacts.pdf",
                ),
                show=show,
            )
        )

        outputs[f"{prefix}_{safe_baseline}"] = (
            plot_loss_vs_propagation_penalty_single_method(
                all_dataset_results,
                method=baseline_method,
                loss_metric=loss_metric,
                output_path=os.path.join(
                    output_dir,
                    f"{prefix}_vs_propagation_penalty_"
                    f"{safe_baseline}.pdf",
                ),
                show=show,
            )
        )

    return outputs



# ============================================================
# ProACTS multiple-position success-rate plot
# ============================================================

def plot_proacts_multiple_success_rate_by_dataset(
    all_dataset_results: Mapping,
    *,
    success_type: str = "Flip Success",
    tolerance: float = 1e-8,
    output_path: Optional[str] = (
        "summary_evaluation/"
        "proacts_multiple_success_rate.pdf"
    ),
    show: bool = True,
) -> pd.DataFrame:
    """
    Plot ProACTS success rate against the number of desired
    positions for every dataset.

    Each factual instance is first aggregated across seeds
    using the median loss.

    success_type:
        "Flip Success"
            uses margin_loss_flipped

        "Margin Success"
            uses margin_loss

    x-axis:
        number of desired positions

    y-axis:
        success rate

    Layout:
        up to 10 datasets -> 2 rows x 5 columns
    """

    if success_type == "Flip Success":

        loss_column = (
            "margin_loss_flipped"
        )

    elif success_type == "Margin Success":

        loss_column = (
            "margin_loss"
        )

    else:

        raise ValueError(
            "success_type must be either "
            "'Flip Success' or 'Margin Success'."
        )

    datasets = list(
        all_dataset_results.keys()
    )

    if not datasets:

        raise ValueError(
            "No dataset results supplied."
        )

    n_datasets = len(
        datasets
    )

    n_cols = min(
        5,
        n_datasets,
    )

    n_rows = math.ceil(
        n_datasets
        /
        n_cols
    )

    fig, axes = plt.subplots(
        n_rows,
        n_cols,
        figsize=(
            3.5 * n_cols,
            3.0 * n_rows,
        ),
        squeeze=False,
        sharey=True,
    )

    axes = axes.flatten()

    summary_rows = []

    for panel_idx, dataset in enumerate(
        datasets
    ):

        ax = axes[
            panel_idx
        ]

        combined = (
            _combine_proacts_multiple_across_seeds(
                all_dataset_results[
                    dataset
                ]
            )
        )

        # ====================================================
        # Aggregate loss per factual instance across seeds
        # ====================================================

        aggregated = (
            _aggregate_proacts_multiple_metric(
                combined,
                loss_column,
            )
        )

        # ====================================================
        # Success rate by number of desired positions
        # ====================================================

        plot_rows = []

        for (
            desired_count,
            group,
        ) in aggregated.groupby(
            DESIRED_COUNT_COLUMN,
            sort=True,
        ):

            values = (
                group[
                    "Value"
                ]
                .to_numpy(
                    dtype=float
                )
            )

            success = (
                values
                <= tolerance
            )

            rate = float(
                success.mean()
            )

            n = int(
                len(
                    success
                )
            )

            plot_rows.append(
                {
                    "Desired Positions": int(
                        desired_count
                    ),
                    "Success Rate": rate,
                    "N": n,
                }
            )

            summary_rows.append(
                {
                    "Dataset": dataset,
                    "Success Type": (
                        success_type
                    ),
                    "Desired Positions": int(
                        desired_count
                    ),
                    "N Instances": n,
                    "Successful": int(
                        success.sum()
                    ),
                    "Success Rate": rate,
                }
            )

        plot_df = pd.DataFrame(
            plot_rows
        )

        # ====================================================
        # Line
        # ====================================================

        ax.plot(
            plot_df[
                "Desired Positions"
            ],
            plot_df[
                "Success Rate"
            ],
            marker="o",
            linewidth=1.8,
            markersize=5,
        )

        ax.set_title(
            dataset
        )

        ax.set_ylim(
            0,
            1.02,
        )

        ax.set_xticks(
            sorted(
                plot_df[
                    "Desired Positions"
                ].unique()
            )
        )

        ax.grid(
            alpha=0.20
        )

    # ========================================================
    # Hide unused panels
    # ========================================================

    for idx in range(
        n_datasets,
        len(
            axes
        ),
    ):

        axes[
            idx
        ].set_visible(
            False
        )

    fig.supxlabel(
        "Number of desired positions"
    )

    fig.supylabel(
        f"{success_type} rate"
    )

    fig.tight_layout()

    # ========================================================
    # Save
    # ========================================================

    if output_path is not None:

        directory = os.path.dirname(
            output_path
        )

        if directory:

            os.makedirs(
                directory,
                exist_ok=True,
            )

        fig.savefig(
            output_path,
            bbox_inches="tight",
        )

    if show:

        plt.show()

    else:

        plt.close(
            fig
        )

    return pd.DataFrame(
        summary_rows
    )


# ============================================================
# ProACTS multiple-position metric boxplots
# ============================================================

def plot_proacts_multiple_metric_boxplots_by_dataset(
    all_dataset_results: Mapping,
    *,
    metric: str = "Propagation Penalty",
    output_path: Optional[str] = None,
    show: bool = True,
    random_seed: int = 42,
) -> pd.DataFrame:
    """
    Publication-style boxplots for ProACTS multiple-position
    experiments.

    Supported metrics:
        - "Margin Loss"
        - "Propagation Penalty"
        - "Distance"
        - "Sparsity"

    Aggregation:
        Each factual instance is first aggregated across
        seeds using the median.

    Plot:
        - filled box = distribution across factual instances
        - jittered dots = individual factual instances
        - black diamond = median
        - one compact n-line per panel = sample sizes from
          left to right

    Encoding:
        x-axis = number of desired positions
        y-axis = selected metric
    """

    # ========================================================
    # Validation
    # ========================================================

    supported_metrics = {
        "Margin Loss",
        "Propagation Penalty",
        "Distance",
        "Sparsity",
    }

    if metric not in supported_metrics:

        raise ValueError(
            f"Unknown metric '{metric}'. "
            f"Expected one of: "
            f"{sorted(supported_metrics)}"
        )

    if metric not in PROACTS_MULTIPLE_METRICS:

        raise KeyError(
            f"Metric '{metric}' is not defined in "
            "PROACTS_MULTIPLE_METRICS."
        )

    source_column = (
        PROACTS_MULTIPLE_METRICS[
            metric
        ]
    )

    rng = np.random.default_rng(
        random_seed
    )

    datasets = list(
        all_dataset_results.keys()
    )

    if not datasets:

        raise ValueError(
            "No dataset results supplied."
        )

    # ========================================================
    # Aggregate each factual instance across seeds
    # ========================================================

    aggregated_by_dataset = {}

    for dataset in datasets:

        combined = (
            _combine_proacts_multiple_across_seeds(
                all_dataset_results[
                    dataset
                ]
            )
        )

        aggregated = (
            _aggregate_proacts_multiple_metric(
                combined,
                source_column,
            )
        )

        aggregated_by_dataset[
            dataset
        ] = aggregated

    # ========================================================
    # Publication-friendly colours
    # ========================================================

    # Okabe-Ito blue
    box_color = "#0072B2"

    # Darker blue for jittered observations
    point_color = "#00394D"

    # ========================================================
    # Figure layout
    # ========================================================

    n_datasets = len(
        datasets
    )

    n_cols = min(
        5,
        n_datasets,
    )

    n_rows = math.ceil(
        n_datasets
        /
        n_cols
    )

    fig, axes = plt.subplots(
        n_rows,
        n_cols,
        figsize=(
            3.6 * n_cols,
            3.2 * n_rows,
        ),
        squeeze=False,
        sharey=False,
    )

    axes = axes.flatten()

    summary_rows = []

    # ========================================================
    # Dataset panels
    # ========================================================

    for panel_idx, dataset in enumerate(
        datasets
    ):

        ax = axes[
            panel_idx
        ]

        aggregated = (
            aggregated_by_dataset[
                dataset
            ]
        )

        # ----------------------------------------------------
        # Dataset has no observations
        # ----------------------------------------------------

        if aggregated.empty:

            ax.set_title(
                dataset,
                fontsize=12,
                pad=5,
            )

            ax.text(
                0.5,
                0.5,
                "No observations",
                ha="center",
                va="center",
                transform=ax.transAxes,
                fontsize=10,
            )

            ax.set_axis_off()

            continue

        desired_counts = sorted(
            aggregated[
                DESIRED_COUNT_COLUMN
            ]
            .dropna()
            .astype(int)
            .unique()
        )

        # ====================================================
        # Build box groups
        # ========================================================

        box_values = []
        plotted_counts = []

        for desired_count in (
            desired_counts
        ):

            group = (
                aggregated.loc[
                    aggregated[
                        DESIRED_COUNT_COLUMN
                    ].eq(
                        desired_count
                    )
                ]
            )

            values = (
                group[
                    "Value"
                ]
                .dropna()
                .to_numpy(
                    dtype=float
                )
            )

            if len(
                values
            ) == 0:

                continue

            box_values.append(
                values
            )

            plotted_counts.append(
                int(
                    desired_count
                )
            )

            median, q1, q3 = (
                _median_iqr_values(
                    values
                )
            )

            summary_rows.append(
                {
                    "Dataset": dataset,

                    "Metric": metric,

                    "Desired Positions": int(
                        desired_count
                    ),

                    "N Instances": int(
                        len(
                            values
                        )
                    ),

                    "Median": median,

                    "Q1": q1,

                    "Q3": q3,

                    "Mean": float(
                        np.mean(
                            values
                        )
                    ),
                }
            )

        # ----------------------------------------------------
        # No finite observations
        # ----------------------------------------------------

        if not box_values:

            ax.set_title(
                dataset,
                fontsize=12,
                pad=5,
            )

            ax.text(
                0.5,
                0.5,
                "No observations",
                ha="center",
                va="center",
                transform=ax.transAxes,
                fontsize=10,
            )

            ax.set_axis_off()

            continue

        positions = np.arange(
            1,
            len(
                box_values
            )
            + 1,
        )

        # ====================================================
        # Boxplots
        # ========================================================

        boxplot = ax.boxplot(
            box_values,

            positions=positions,

            widths=0.62,

            patch_artist=True,

            # Jittered points already show all observations.
            showfliers=False,

            medianprops={
                "color": "black",
                "linewidth": 1.7,
            },

            whiskerprops={
                "color": "black",
                "linewidth": 1.0,
            },

            capprops={
                "color": "black",
                "linewidth": 1.0,
            },

            boxprops={
                "edgecolor": "black",
                "linewidth": 1.0,
            },
        )

        # ====================================================
        # Fill boxes
        # ========================================================

        for patch in (
            boxplot[
                "boxes"
            ]
        ):

            patch.set_facecolor(
                box_color
            )

            patch.set_alpha(
                0.65
            )

        # ====================================================
        # Individual observations + median marker
        # ========================================================

        for (
            position,
            values,
        ) in zip(
            positions,
            box_values,
        ):

            # ------------------------------------------------
            # Horizontal jitter
            # ------------------------------------------------

            jitter = rng.normal(
                loc=0.0,
                scale=0.055,
                size=len(
                    values
                ),
            )

            # ------------------------------------------------
            # Individual factual instances
            # ------------------------------------------------

            ax.scatter(
                position
                +
                jitter,

                values,

                s=15,

                color=point_color,

                alpha=0.32,

                edgecolors="none",

                zorder=2,
            )

            # ------------------------------------------------
            # Median diamond
            # ------------------------------------------------

            median_value = float(
                np.median(
                    values
                )
            )

            ax.scatter(
                position,

                median_value,

                marker="D",

                s=35,

                facecolor="black",

                edgecolor="black",

                linewidth=0.5,

                zorder=4,
            )

        # ====================================================
        # Y-range
        # ========================================================

        all_values = np.concatenate(
            box_values
        )

        value_min = float(
            np.min(
                all_values
            )
        )

        value_max = float(
            np.max(
                all_values
            )
        )

        value_range = (
            value_max
            -
            value_min
        )

        if np.isclose(
            value_range,
            0.0,
        ):

            value_range = max(
                abs(
                    value_max
                ),
                1.0,
            )

        # ====================================================
        # Compact sample-size annotation
        #
        # Example:
        # n = 20; 18; 15; 12
        #
        # Values correspond left-to-right to x-axis positions.
        # ========================================================

        sample_sizes = [
            len(values)
            for values in box_values
        ]

        sample_size_text = (
            r"$n$ = "
            +
            "; ".join(
                str(n)
                for n in sample_sizes
            )
        )

        ax.text(
            0.98,
            0.96,

            sample_size_text,

            transform=ax.transAxes,

            ha="right",
            va="top",

            fontsize=8.5,
        )

        # ====================================================
        # X axis
        # ========================================================

        ax.set_xticks(
            positions
        )

        ax.set_xticklabels(
            [
                str(
                    desired_count
                )
                for desired_count
                in plotted_counts
            ]
        )

        # ====================================================
        # Dataset title
        # ========================================================

        ax.set_title(
            dataset,

            fontsize=12,

            pad=5,
        )

        # ====================================================
        # Y limits
        # ========================================================

        lower_limit = min(
            0.0,
            value_min,
        )

        # No longer need extra vertical space for n= above
        # each individual box.
        upper_limit = (
            value_max
            +
            0.10
            *
            value_range
        )

        ax.set_ylim(
            lower_limit,
            upper_limit,
        )

        # ====================================================
        # Grid
        # ========================================================

        ax.grid(
            axis="y",
            linestyle="--",
            linewidth=0.6,
            alpha=0.25,
        )

        ax.set_axisbelow(
            True
        )

        # ====================================================
        # Publication appearance
        # ========================================================

        ax.spines[
            "top"
        ].set_visible(
            False
        )

        ax.spines[
            "right"
        ].set_visible(
            False
        )

        # Tick labels / numbers = 9 pt
        ax.tick_params(
            axis="both",
            labelsize=9,
            length=4,
            width=0.8,
        )

    # ========================================================
    # Hide unused panels
    # ========================================================

    for idx in range(
        n_datasets,
        len(
            axes
        ),
    ):

        axes[
            idx
        ].set_visible(
            False
        )

    # ========================================================
    # Shared axis titles = 14 pt
    # ========================================================

    fig.supxlabel(
        "Number of desired positions",
        fontsize=14,
    )

    fig.supylabel(
        metric,
        fontsize=14,
    )

    # ========================================================
    # Median legend
    # ========================================================

    median_handle = plt.Line2D(
        [0],
        [0],

        marker="D",

        linestyle="",

        markersize=7,

        markerfacecolor="black",

        markeredgecolor="black",

        label="Median",
    )

    fig.legend(
        handles=[
            median_handle
        ],

        loc="upper center",

        bbox_to_anchor=(
            0.5,
            0.995,
        ),

        frameon=False,

        # Legend = 12 pt
        fontsize=12,
    )

    # ========================================================
    # Manual spacing
    # ========================================================

    fig.subplots_adjust(
        left=0.065,
        right=0.985,
        bottom=0.10,
        top=0.90,
        wspace=0.16,
        hspace=0.24,
    )

    # ========================================================
    # Save
    # ========================================================

    if output_path is not None:

        directory = os.path.dirname(
            output_path
        )

        if directory:

            os.makedirs(
                directory,
                exist_ok=True,
            )

        fig.savefig(
            output_path,
            bbox_inches="tight",
        )

    # ========================================================
    # Show / close
    # ========================================================

    if show:

        plt.show()

    else:

        plt.close(
            fig
        )

    return pd.DataFrame(
        summary_rows
    )


# ============================================================
# Generate all ProACTS multiple-position plots
# ============================================================

def generate_proacts_multiple_position_plots(
    all_dataset_results: Mapping,
    *,
    output_dir: str = (
        "summary_evaluation"
    ),
    tolerance: float = 1e-8,
    show: bool = True,
) -> dict:
    """
    Generate the main multiple-position figures for ProACTS.
    """

    os.makedirs(
        output_dir,
        exist_ok=True,
    )

    outputs = {}

    # ========================================================
    # Success
    # ========================================================

    outputs[
        "flip_success"
    ] = (
        plot_proacts_multiple_success_rate_by_dataset(
            all_dataset_results,
            success_type="Flip Success",
            tolerance=tolerance,
            output_path=os.path.join(
                output_dir,
                "proacts_multiple_flip_success.pdf",
            ),
            show=show,
        )
    )

    outputs[
        "margin_success"
    ] = (
        plot_proacts_multiple_success_rate_by_dataset(
            all_dataset_results,
            success_type="Margin Success",
            tolerance=tolerance,
            output_path=os.path.join(
                output_dir,
                "proacts_multiple_margin_success.pdf",
            ),
            show=show,
        )
    )

    # ========================================================
    # Continuous metrics
    # ========================================================

    for metric, filename in [
        (
            "Propagation Penalty",
            "proacts_multiple_propagation_penalty.pdf",
        ),
        (
            "Margin Loss",
            "proacts_multiple_margin_loss.pdf",
        ),
        (
            "Distance",
            "proacts_multiple_distance.pdf",
        ),
        (
            "Sparsity",
            "proacts_multiple_sparsity.pdf",
        ),
    ]:

        outputs[
            metric
        ] = (
            plot_proacts_multiple_metric_boxplots_by_dataset(
                all_dataset_results,
                metric=metric,
                output_path=os.path.join(
                    output_dir,
                    filename,
                ),
                show=show,
            )
        )

    return outputs


# ============================================================
# ProACTS multiple-position metric heatmap
# ============================================================

def plot_proacts_multiple_metric_heatmap(
    all_dataset_results: Mapping,
    *,
    metric: str = "Propagation Penalty",
    aggregation: str = "median",
    output_path: Optional[str] = None,
    show: bool = True,
) -> pd.DataFrame:
    """
    Heatmap for one ProACTS multiple-position metric.

    Rows:
        datasets

    Columns:
        number of desired positions

    Cell:
        aggregate metric value across factual instances

    Aggregation:
        1. Each factual instance is first aggregated across
           seeds using the median via
           _aggregate_proacts_multiple_metric().

        2. The resulting per-instance median values are then
           aggregated within each dataset x desired-position
           count using either the median or mean.

    Missing dataset x desired-position combinations:
        shown as white cells with "--".

    Supported aggregation:
        "median"
        "mean"
    """

    # ========================================================
    # Validation
    # ========================================================

    if metric not in PROACTS_MULTIPLE_METRICS:

        raise ValueError(
            f"Unknown metric '{metric}'. "
            f"Expected one of: "
            f"{list(PROACTS_MULTIPLE_METRICS.keys())}"
        )

    if aggregation not in {
        "median",
        "mean",
    }:

        raise ValueError(
            "aggregation must be either "
            "'median' or 'mean'."
        )

    source_column = (
        PROACTS_MULTIPLE_METRICS[
            metric
        ]
    )

    rows = []

    # ========================================================
    # Build summary values
    # ========================================================

    for (
        dataset,
        seed_results,
    ) in all_dataset_results.items():

        combined = (
            _combine_proacts_multiple_across_seeds(
                seed_results
            )
        )

        # ----------------------------------------------------
        # Per factual instance:
        # median metric across seeds
        # ----------------------------------------------------

        aggregated = (
            _aggregate_proacts_multiple_metric(
                combined,
                source_column,
            )
        )

        if aggregated.empty:
            continue

        # ----------------------------------------------------
        # Aggregate the resulting per-instance values
        # by number of desired positions
        # ----------------------------------------------------

        for (
            desired_count,
            group,
        ) in aggregated.groupby(
            DESIRED_COUNT_COLUMN,
            sort=True,
        ):

            values = (
                group[
                    "Value"
                ]
                .dropna()
                .to_numpy(
                    dtype=float
                )
            )

            if len(
                values
            ) == 0:
                continue

            # -----------------------------------------------
            # Heatmap cell value
            # -----------------------------------------------

            if aggregation == "median":

                cell_value = float(
                    np.median(
                        values
                    )
                )

            else:

                cell_value = float(
                    np.mean(
                        values
                    )
                )

            # -----------------------------------------------
            # Additional descriptive statistics
            # -----------------------------------------------

            q1, median, q3 = np.quantile(
                values,
                [
                    0.25,
                    0.50,
                    0.75,
                ],
            )

            rows.append(
                {
                    "Dataset": dataset,

                    "Metric": metric,

                    "Desired Positions": int(
                        desired_count
                    ),

                    "N Instances": int(
                        len(
                            values
                        )
                    ),

                    "Value": (
                        cell_value
                    ),

                    "Median": float(
                        median
                    ),

                    "Q1": float(
                        q1
                    ),

                    "Q3": float(
                        q3
                    ),

                    "Mean": float(
                        np.mean(
                            values
                        )
                    ),
                }
            )

    # ========================================================
    # Summary DataFrame
    # ========================================================

    summary = pd.DataFrame(
        rows
    )

    if summary.empty:

        raise ValueError(
            "No multiple-position observations "
            "were available."
        )

    # ========================================================
    # Heatmap matrix
    #
    # Missing combinations automatically become NaN.
    # ========================================================

    matrix = (
        summary.pivot(
            index="Dataset",
            columns="Desired Positions",
            values="Value",
        )
    )

    # --------------------------------------------------------
    # Preserve original dataset order
    # --------------------------------------------------------

    dataset_order = [
        dataset
        for dataset
        in all_dataset_results.keys()
        if dataset in matrix.index
    ]

    matrix = matrix.reindex(
        dataset_order
    )

    # --------------------------------------------------------
    # Sort desired-position counts
    # --------------------------------------------------------

    matrix = matrix.reindex(
        sorted(
            matrix.columns
        ),
        axis=1,
    )

    # ========================================================
    # Values
    # ========================================================

    values = matrix.to_numpy(
        dtype=float
    )

    # --------------------------------------------------------
    # Mask missing combinations.
    #
    # NaN = no experiments for this dataset x desired-count.
    # --------------------------------------------------------

    masked_values = (
        np.ma.masked_invalid(
            values
        )
    )

    # ========================================================
    # Figure dimensions
    # ========================================================

    fig_width = max(
        5.0,
        1.1
        *
        len(
            matrix.columns
        )
        +
        3.0,
    )

    fig_height = max(
        4.0,
        0.45
        *
        len(
            matrix.index
        )
        +
        1.8,
    )

    fig, ax = plt.subplots(
        figsize=(
            fig_width,
            fig_height,
        )
    )

    # ========================================================
    # Colour map
    # ========================================================

    cmap = (
        plt.get_cmap(
            "cividis"
        )
        .copy()
    )

    # Missing values:
    # white = no experiment exists.
    cmap.set_bad(
        color="white"
    )

    # ========================================================
    # Heatmap
    # ========================================================

    image = ax.imshow(
        masked_values,
        aspect="auto",
        interpolation="nearest",
        cmap=cmap,
    )

    # ========================================================
    # Axes
    # ========================================================

    ax.set_xticks(
        np.arange(
            len(
                matrix.columns
            )
        )
    )

    ax.set_xticklabels(
        [
            str(
                int(
                    desired_count
                )
            )
            for desired_count
            in matrix.columns
        ]
    )

    ax.set_yticks(
        np.arange(
            len(
                matrix.index
            )
        )
    )

    ax.set_yticklabels(
        matrix.index
    )

    ax.set_xlabel(
        "Number of desired positions"
    )

    ax.set_ylabel(
        "Dataset"
    )

    # I would omit the title in the paper if the figure
    # caption already states the metric.
    ax.set_title(
        metric
    )

    # ========================================================
    # Determine midpoint only from observed cells
    # ========================================================

    finite_values = (
        values[
            np.isfinite(
                values
            )
        ]
    )

    if len(
        finite_values
    ) > 0:

        value_midpoint = (
            float(
                np.min(
                    finite_values
                )
            )
            +
            float(
                np.max(
                    finite_values
                )
            )
        ) / 2.0

    else:

        value_midpoint = 0.0

    # ========================================================
    # Cell annotations
    # ========================================================

    for row_idx in range(
        values.shape[
            0
        ]
    ):

        for col_idx in range(
            values.shape[
                1
            ]
        ):

            value = values[
                row_idx,
                col_idx,
            ]

            # ------------------------------------------------
            # No experiment
            # ------------------------------------------------

            if not np.isfinite(
                value
            ):

                ax.text(
                    col_idx,
                    row_idx,
                    "--",
                    ha="center",
                    va="center",
                    fontsize=8,
                )

                continue

            # ------------------------------------------------
            # Observed value
            # ------------------------------------------------

            text_color = (
                "white"
                if value > value_midpoint
                else "black"
            )

            ax.text(
                col_idx,
                row_idx,
                f"{value:.3f}",
                ha="center",
                va="center",
                fontsize=8,
                color=text_color,
            )

    # ========================================================
    # Cell boundaries
    #
    # Makes missing white cells easier to distinguish.
    # ========================================================

    ax.set_xticks(
        np.arange(
            -0.5,
            len(
                matrix.columns
            ),
            1,
        ),
        minor=True,
    )

    ax.set_yticks(
        np.arange(
            -0.5,
            len(
                matrix.index
            ),
            1,
        ),
        minor=True,
    )

    ax.grid(
        which="minor",
        linewidth=0.5,
    )

    ax.tick_params(
        which="minor",
        bottom=False,
        left=False,
    )

    # ========================================================
    # Colorbar
    # ========================================================

    colorbar = fig.colorbar(
        image,
        ax=ax,
        fraction=0.035,
        pad=0.03,
    )

    colorbar.set_label(
        metric
    )

    # ========================================================
    # Layout
    # ========================================================

    fig.tight_layout()

    # ========================================================
    # Save
    # ========================================================

    if output_path is not None:

        directory = os.path.dirname(
            output_path
        )

        if directory:

            os.makedirs(
                directory,
                exist_ok=True,
            )

        fig.savefig(
            output_path,
            bbox_inches="tight",
        )

    # ========================================================
    # Show / close
    # ========================================================

    if show:

        plt.show()

    else:

        plt.close(
            fig
        )

    return summary


# ============================================================
# ProACTS multiple:
# Margin loss vs propagation penalty
# Colored by number of desired positions
# ============================================================

def plot_proacts_multiple_margin_vs_propagation(
    all_dataset_results: Mapping,
    *,
    output_path: Optional[str] = (
        "summary_evaluation/"
        "multiple_margin_vs_propagation.pdf"
    ),
    show: bool = True,
) -> pd.DataFrame:
    """
    Scatter plot of margin loss versus propagation penalty
    for ProACTS multiple-position counterfactuals.

    Each point:
        one factual instance.

    Before plotting:
        margin loss and propagation penalty are each
        aggregated across seeds using the median.

    x-axis:
        Margin loss

    y-axis:
        Propagation penalty

    color:
        Number of desired positions

    Layout:
        up to 10 datasets -> 2 rows x 5 columns.
    """

    datasets = list(
        all_dataset_results.keys()
    )

    if not datasets:
        raise ValueError(
            "No dataset results supplied."
        )

    # ========================================================
    # First build all plotting data
    # ========================================================

    plot_frames = []

    for dataset, seed_results in (
        all_dataset_results.items()
    ):

        combined = (
            _combine_proacts_multiple_across_seeds(
                seed_results
            )
        )

        # ----------------------------------------------------
        # Median margin loss per factual instance across seeds
        # ----------------------------------------------------

        margin = (
            _aggregate_proacts_multiple_metric(
                combined,
                "margin_loss",
            )
            .rename(
                columns={
                    "Value": "MarginLoss",
                }
            )
        )

        # ----------------------------------------------------
        # Median propagation penalty per factual instance
        # across seeds
        # ----------------------------------------------------

        propagation = (
            _aggregate_proacts_multiple_metric(
                combined,
                "process_violation",
            )
            .rename(
                columns={
                    "Value": "PropagationPenalty",
                }
            )
        )

        # ----------------------------------------------------
        # Pair the two metrics for the same factual query
        # and desired-position count.
        # ----------------------------------------------------

        frame = (
            margin[
                [
                    *ID_COLUMNS,
                    DESIRED_COUNT_COLUMN,
                    "MarginLoss",
                ]
            ]
            .merge(
                propagation[
                    [
                        *ID_COLUMNS,
                        DESIRED_COUNT_COLUMN,
                        "PropagationPenalty",
                    ]
                ],
                on=[
                    *ID_COLUMNS,
                    DESIRED_COUNT_COLUMN,
                ],
                how="inner",
                validate="one_to_one",
            )
        )

        frame.insert(
            0,
            "Dataset",
            dataset,
        )

        plot_frames.append(
            frame
        )

    plot_data = pd.concat(
        plot_frames,
        ignore_index=True,
    )

    if plot_data.empty:
        raise ValueError(
            "No multiple-position observations "
            "available."
        )

    # ========================================================
    # Desired-position range
    # ========================================================

    desired_counts = sorted(
        plot_data[
            DESIRED_COUNT_COLUMN
        ]
        .dropna()
        .astype(int)
        .unique()
    )

    desired_min = min(
        desired_counts
    )

    desired_max = max(
        desired_counts
    )

    # ========================================================
    # Layout
    # ========================================================

    n_datasets = len(
        datasets
    )

    n_cols = min(
        5,
        n_datasets,
    )

    n_rows = math.ceil(
        n_datasets
        /
        n_cols
    )

    fig, axes = plt.subplots(
        n_rows,
        n_cols,
        figsize=(
            3.5 * n_cols,
            3.0 * n_rows,
        ),
        squeeze=False,
    )

    axes = axes.flatten()

    # ========================================================
    # Discrete desired-count color mapping
    # ========================================================

    cmap = plt.get_cmap(
        "cividis",
        len(
            desired_counts
        ),
    )

    # Boundaries make integer desired counts map cleanly
    # to discrete colours.
    boundaries = np.arange(
        desired_min - 0.5,
        desired_max + 1.5,
        1,
    )

    norm = plt.matplotlib.colors.BoundaryNorm(
        boundaries,
        cmap.N,
    )

    scatter_for_colorbar = None

    # ========================================================
    # Dataset panels
    # ========================================================

    for panel_idx, dataset in enumerate(
        datasets
    ):

        ax = axes[
            panel_idx
        ]

        dataset_df = (
            plot_data.loc[
                plot_data[
                    "Dataset"
                ].eq(
                    dataset
                )
            ]
            .copy()
        )

        if dataset_df.empty:

            ax.set_title(
                dataset
            )

            ax.text(
                0.5,
                0.5,
                "No observations",
                ha="center",
                va="center",
                transform=ax.transAxes,
            )

            ax.set_axis_off()

            continue

        x = (
            dataset_df[
                "MarginLoss"
            ]
            .to_numpy(
                dtype=float
            )
        )

        y = (
            dataset_df[
                "PropagationPenalty"
            ]
            .to_numpy(
                dtype=float
            )
        )

        desired = (
            dataset_df[
                DESIRED_COUNT_COLUMN
            ]
            .to_numpy(
                dtype=int
            )
        )

        scatter = ax.scatter(
            x,
            y,
            c=desired,
            cmap=cmap,
            norm=norm,
            s=24,
            alpha=0.60,
            edgecolors="none",
        )

        scatter_for_colorbar = (
            scatter
        )

        ax.set_title(
            dataset
        )

        ax.set_xlim(
            left=0
        )

        ax.set_ylim(
            bottom=0
        )

        ax.grid(
            alpha=0.20
        )

    # ========================================================
    # Hide unused panels
    # ========================================================

    for idx in range(
        n_datasets,
        len(
            axes
        ),
    ):

        axes[
            idx
        ].set_visible(
            False
        )

    # ========================================================
    # Shared labels
    # ========================================================

    fig.supxlabel(
        "Margin loss"
    )

    fig.supylabel(
        "Propagation penalty"
    )

    # ========================================================
    # Shared colorbar
    # ========================================================

    if scatter_for_colorbar is not None:

        visible_axes = [
            ax
            for ax in axes
            if ax.get_visible()
        ]

        colorbar = fig.colorbar(
            scatter_for_colorbar,
            ax=visible_axes,
            fraction=0.025,
            pad=0.02,
            ticks=desired_counts,
        )

        colorbar.set_label(
            "Number of desired positions"
        )

    # Leave room for colorbar.
    fig.subplots_adjust(
        left=0.06,
        right=0.90,
        bottom=0.10,
        top=0.94,
        wspace=0.28,
        hspace=0.35,
    )

    # ========================================================
    # Save
    # ========================================================

    if output_path is not None:

        directory = os.path.dirname(
            output_path
        )

        if directory:

            os.makedirs(
                directory,
                exist_ok=True,
            )

        fig.savefig(
            output_path,
            bbox_inches="tight",
        )

    if show:

        plt.show()

    else:

        plt.close(
            fig
        )

    return plot_data



# ============================================================
# Runtime summary
# ============================================================

def runtime_table_all_seeds(
    runtime_paths: dict,
    all_dataset_results: dict,
    *,
    output_path: str | None = None,
) -> pd.DataFrame:
    """
    Create runtime summary across all seeds.

    Expected runtime_paths format
    -----------------------------
    {
        "dataset_1": {
            42: {
                "ProACTS": "path/to/proacts_seed42.log",
                "DiCE": "path/to/dice_seed42.log",
            },
            43: {
                "ProACTS": "path/to/proacts_seed43.log",
                "DiCE": "path/to/dice_seed43.log",
            },
        },
        ...
    }

    For each dataset and seed:

        N = number of single-position ProACTS experiments
            in all_dataset_results[dataset][seed]["ProACTS"]

        ProACTS:
            first N runtime entries are used.

        DiCE:
            exactly 2N entries are expected.

            first N  -> DiCE4EL
            next N   -> DiCE4EL+

    Runtime observations from all seeds are then pooled.

    The main reported statistic is:

        Median [Q1, Q3]

    Mean ± SD is also returned as a diagnostic.

    All runtimes are optimization-loop runtimes in seconds
    per factual query.
    """

    # ========================================================
    # Runtime patterns
    # ========================================================

    proacts_pattern = (
        r"Time taken GA loop:\s*"
        r"([0-9]*\.?[0-9]+)\s*seconds"
    )

    dice_pattern = (
        r"Time taken for DiCE4EL optimization loop:\s*"
        r"([0-9]*\.?[0-9]+)\s*seconds"
    )

    # ========================================================
    # Extract runtimes
    # ========================================================

    def extract(
        log_path: str,
        pattern: str,
    ) -> np.ndarray:

        with open(
            log_path,
            "r",
            encoding="utf-8",
            errors="ignore",
        ) as f:
            text = f.read()

        values = re.findall(
            pattern,
            text,
        )

        if not values:
            raise ValueError(
                "No runtime entries found in:\n"
                f"{log_path}"
            )

        return np.asarray(
            values,
            dtype=float,
        )

    # ========================================================
    # Formatting
    # ========================================================

    def median_iqr(
        values: np.ndarray,
    ) -> tuple[float, float, float]:

        q1, median, q3 = np.quantile(
            values,
            [
                0.25,
                0.50,
                0.75,
            ],
        )

        return (
            float(median),
            float(q1),
            float(q3),
        )

    def median_iqr_text(
        values: np.ndarray,
    ) -> str:

        median, q1, q3 = median_iqr(
            values
        )

        return (
            f"{median:.3f} "
            f"[{q1:.3f}, {q3:.3f}]"
        )

    def mean_sd_text(
        values: np.ndarray,
    ) -> str:

        mean = float(
            np.mean(values)
        )

        sd = (
            float(
                np.std(
                    values,
                    ddof=1,
                )
            )
            if len(values) > 1
            else np.nan
        )

        if np.isfinite(sd):

            return (
                f"{mean:.3f} "
                f"± {sd:.3f}"
            )

        return f"{mean:.3f}"

    # ========================================================
    # Process datasets
    # ========================================================

    rows = []

    for dataset, seed_paths in (
        runtime_paths.items()
    ):

        if dataset not in all_dataset_results:

            raise KeyError(
                f"Dataset '{dataset}' not found "
                "in all_dataset_results."
            )

        # ----------------------------------------------------
        # Store every query runtime across every seed
        # ----------------------------------------------------

        all_proacts = []
        all_dice4el = []
        all_dice4el_plus = []

        seeds_used = []

        # ====================================================
        # Seeds
        # ====================================================

        for seed, paths in seed_paths.items():

            if seed not in all_dataset_results[
                dataset
            ]:

                raise KeyError(
                    f"Seed {seed} not found for "
                    f"dataset '{dataset}'."
                )

            seed_results = (
                all_dataset_results[
                    dataset
                ][seed]
            )

            if "ProACTS" not in seed_results:

                raise KeyError(
                    f"Dataset '{dataset}', "
                    f"seed {seed} does not contain "
                    "ProACTS results."
                )

            # ------------------------------------------------
            # Number of SINGLE-position queries
            # ------------------------------------------------

            n_single = int(
                len(
                    seed_results[
                        "ProACTS"
                    ]
                )
            )

            if n_single <= 0:

                raise ValueError(
                    f"{dataset}, seed {seed}: "
                    "number of single-position "
                    "experiments is zero."
                )

            # =================================================
            # ProACTS
            # =================================================

            proacts_all = extract(
                paths["ProACTS"],
                proacts_pattern,
            )

            if len(proacts_all) < n_single:

                raise ValueError(
                    f"{dataset}, seed {seed}: "
                    f"expected at least {n_single} "
                    "ProACTS runtime entries, "
                    f"but found {len(proacts_all)}."
                )

            # First N correspond to single-position runs.
            proacts = (
                proacts_all[
                    :n_single
                ]
            )

            # =================================================
            # DiCE4EL / DiCE4EL+
            # =================================================

            dice_all = extract(
                paths["DiCE"],
                dice_pattern,
            )

            expected_dice = (
                2 * n_single
            )

            if len(dice_all) != expected_dice:

                raise ValueError(
                    f"{dataset}, seed {seed}: "
                    f"expected exactly "
                    f"{expected_dice} DiCE runtime "
                    f"entries (2 × {n_single}), "
                    f"but found {len(dice_all)}."
                )

            # First N
            dice4el = (
                dice_all[
                    :n_single
                ]
            )

            # Second N
            dice4el_plus = (
                dice_all[
                    n_single:
                    2 * n_single
                ]
            )

            # =================================================
            # Sanity check
            # =================================================

            if not (
                len(proacts)
                ==
                len(dice4el)
                ==
                len(dice4el_plus)
                ==
                n_single
            ):

                raise RuntimeError(
                    f"{dataset}, seed {seed}: "
                    "runtime counts are inconsistent."
                )

            print(
                f"{dataset}, seed={seed}: "
                f"N={n_single}, "
                f"ProACTS={len(proacts)}, "
                f"DiCE4EL={len(dice4el)}, "
                f"DiCE4EL+={len(dice4el_plus)}, "
                f"ProACTS total log="
                f"{len(proacts_all)}"
            )

            # ------------------------------------------------
            # Pool individual query runtimes
            # ------------------------------------------------

            all_proacts.extend(
                proacts.tolist()
            )

            all_dice4el.extend(
                dice4el.tolist()
            )

            all_dice4el_plus.extend(
                dice4el_plus.tolist()
            )

            seeds_used.append(
                seed
            )

        # ====================================================
        # Convert pooled observations to arrays
        # ====================================================

        proacts_values = np.asarray(
            all_proacts,
            dtype=float,
        )

        dice4el_values = np.asarray(
            all_dice4el,
            dtype=float,
        )

        dice4el_plus_values = np.asarray(
            all_dice4el_plus,
            dtype=float,
        )

        # ----------------------------------------------------
        # Final sanity check
        # ----------------------------------------------------

        if not (
            len(proacts_values)
            ==
            len(dice4el_values)
            ==
            len(dice4el_plus_values)
        ):

            raise RuntimeError(
                f"{dataset}: pooled runtime counts "
                "differ between methods."
            )

        # ====================================================
        # Numerical statistics
        # ====================================================

        proacts_med, proacts_q1, proacts_q3 = (
            median_iqr(
                proacts_values
            )
        )

        dice_med, dice_q1, dice_q3 = (
            median_iqr(
                dice4el_values
            )
        )

        dice_plus_med, dice_plus_q1, dice_plus_q3 = (
            median_iqr(
                dice4el_plus_values
            )
        )

        # ====================================================
        # Dataset row
        # ====================================================

        rows.append(
            {
                "Dataset": dataset,

                "N Seeds": int(
                    len(seeds_used)
                ),

                "N Runtime Observations": int(
                    len(proacts_values)
                ),

                # --------------------------------------------
                # Main paper values
                # --------------------------------------------

                "ProACTS Runtime Median [Q1, Q3] (s)":
                    median_iqr_text(
                        proacts_values
                    ),

                "DiCE4EL Runtime Median [Q1, Q3] (s)":
                    median_iqr_text(
                        dice4el_values
                    ),

                "DiCE4EL+ Runtime Median [Q1, Q3] (s)":
                    median_iqr_text(
                        dice4el_plus_values
                    ),

                # --------------------------------------------
                # Raw median values
                # --------------------------------------------

                "ProACTS Median Runtime (s)":
                    proacts_med,

                "DiCE4EL Median Runtime (s)":
                    dice_med,

                "DiCE4EL+ Median Runtime (s)":
                    dice_plus_med,

                # --------------------------------------------
                # Mean ± SD diagnostics
                # --------------------------------------------

                "ProACTS Runtime Mean ± SD (s)":
                    mean_sd_text(
                        proacts_values
                    ),

                "DiCE4EL Runtime Mean ± SD (s)":
                    mean_sd_text(
                        dice4el_values
                    ),

                "DiCE4EL+ Runtime Mean ± SD (s)":
                    mean_sd_text(
                        dice4el_plus_values
                    ),
            }
        )

    # ========================================================
    # Final table
    # ========================================================

    table = pd.DataFrame(
        rows
    )

    numeric_median_columns = [
        "ProACTS Median Runtime (s)",
        "DiCE4EL Median Runtime (s)",
        "DiCE4EL+ Median Runtime (s)",
    ]

    table[
        numeric_median_columns
    ] = (
        table[
            numeric_median_columns
        ]
        .round(3)
    )

    # ========================================================
    # Optional Excel output
    # ========================================================

    if output_path is not None:

        directory = os.path.dirname(
            output_path
        )

        if directory:

            os.makedirs(
                directory,
                exist_ok=True,
            )

        table.to_excel(
            output_path,
            index=False,
        )

    return table



# ============================================================
# Propagation prevalence analysis
# ============================================================

def propagation_prevalence_summary(
    all_dataset_results: Mapping,
    *,
    methods: Sequence[str] = (
        "ProACTS",
        "DiCE4EL+",
    ),
    tolerance: float = 1e-8,
    output_path: Optional[str] = (
        "summary_evaluation/"
        "propagation_prevalence.xlsx"
    ),
) -> pd.DataFrame:
    """
    Quantify propagation prevalence among target-successful
    counterfactuals for ProACTS and baseline methods.

    Experimental unit
    -----------------
    One factual query.

    Each metric is first aggregated across seeds using the
    median via _aggregate_method_metric_across_seeds().

    Success semantics
    -----------------
    ProACTS:
        median margin_loss_flipped <= tolerance
        (autoregressive target success)

    DiCE4EL / DiCE4EL+:
        median margin_loss_na_flipped <= tolerance
        (non-autoregressive target success)

    Propagation prevalence
    ----------------------
    Among target-successful CFs:

        median process_violation > tolerance

    AR target failure
    -----------------
    DiCE4EL / DiCE4EL+:
        CF succeeds under NA evaluation but fails under
        autoregressive evaluation:

        median margin_loss_na_flipped <= tolerance
        AND
        median margin_loss_flipped > tolerance

    ProACTS:
        Target success is already defined autoregressively.
        Therefore, among successful ProACTS CFs,
        AR target failure is necessarily 0.

    Parameters
    ----------
    all_dataset_results:
        Existing all_results dictionary produced by
        evaluate_seed().

    methods:
        Methods to evaluate. Supported:
        "ProACTS", "DiCE4EL", "DiCE4EL+".

    tolerance:
        Numerical tolerance for zero loss.

    output_path:
        Optional Excel output path.

    Returns
    -------
    pd.DataFrame
        One row per dataset and method, plus one pooled
        Overall row per method.
    """

    # ========================================================
    # Validation
    # ========================================================

    valid_methods = {
        "ProACTS",
        "DiCE4EL",
        "DiCE4EL+",
    }

    invalid_methods = (
        set(methods)
        - valid_methods
    )

    if invalid_methods:
        raise ValueError(
            f"Unknown methods: {invalid_methods}. "
            f"Valid methods are: {valid_methods}."
        )

    rows = []

    # ========================================================
    # Datasets
    # ========================================================

    for (
        dataset,
        seed_results,
    ) in all_dataset_results.items():

        # ----------------------------------------------------
        # Combine already-loaded best CFs across seeds
        # ----------------------------------------------------

        combined = (
            _combine_method_results_across_seeds(
                seed_results
            )
        )

        # ====================================================
        # Methods
        # ====================================================

        for method in methods:

            # =================================================
            # Median NA flip loss per factual query
            # =================================================

            na_flip = (
                _aggregate_method_metric_across_seeds(
                    combined=combined,
                    method=method,
                    source_column=(
                        "margin_loss_na_flipped"
                    ),
                )
                .rename(
                    columns={
                        "Value":
                            "NA_Flip_Loss",

                        "N_Seeds":
                            "NA_N_Seeds",
                    }
                )
            )

            # =================================================
            # Median AR flip loss per factual query
            # =================================================

            ar_flip = (
                _aggregate_method_metric_across_seeds(
                    combined=combined,
                    method=method,
                    source_column=(
                        "margin_loss_flipped"
                    ),
                )
                .rename(
                    columns={
                        "Value":
                            "AR_Flip_Loss",

                        "N_Seeds":
                            "AR_N_Seeds",
                    }
                )
            )

            # =================================================
            # Median propagation penalty per factual query
            # =================================================

            propagation = (
                _aggregate_method_metric_across_seeds(
                    combined=combined,
                    method=method,
                    source_column=(
                        "process_violation"
                    ),
                )
                .rename(
                    columns={
                        "Value":
                            "Propagation_Penalty",

                        "N_Seeds":
                            "Propagation_N_Seeds",
                    }
                )
            )

            # =================================================
            # Keep factual queries having all three metrics
            # =================================================

            aggregated = (
                na_flip
                .merge(
                    ar_flip,
                    on=ID_COLUMNS,
                    how="inner",
                    validate="one_to_one",
                )
                .merge(
                    propagation,
                    on=ID_COLUMNS,
                    how="inner",
                    validate="one_to_one",
                )
            )

            n_queries = int(
                len(aggregated)
            )

            # =================================================
            # Method-specific success semantics
            # =================================================

            if method == "ProACTS":

                # ProACTS is optimized/evaluated using
                # autoregressive target behavior.
                success_column = (
                    "AR_Flip_Loss"
                )

                success_semantics = (
                    "AR"
                )

            else:

                # DiCE4EL / DiCE4EL+ define target success
                # using the non-autoregressive prediction.
                success_column = (
                    "NA_Flip_Loss"
                )

                success_semantics = (
                    "NA"
                )

            # =================================================
            # Target-successful factual queries
            # =================================================

            successful = (
                aggregated.loc[
                    aggregated[
                        success_column
                    ]
                    <= tolerance
                ]
                .copy()
            )

            n_successful = int(
                len(successful)
            )

            # =================================================
            # No successful CFs
            # =================================================

            if n_successful == 0:

                rows.append(
                    {
                        "Dataset":
                            dataset,

                        "Method":
                            method,

                        "Success Semantics":
                            success_semantics,

                        "N Factual Queries":
                            n_queries,

                        "Target-Successful CFs":
                            0,

                        "Target Success (%)":
                            0.0,

                        "CFs with Propagation":
                            0,

                        "Propagation Prevalence (%)":
                            np.nan,

                        "AR Target Failures":
                            0,

                        "AR Target Failure (%)":
                            np.nan,

                        "Median Propagation Penalty":
                            np.nan,

                        "Propagation Penalty [Q1, Q3]":
                            "NA",
                    }
                )

                continue

            # =================================================
            # Propagation prevalence
            # =================================================

            has_propagation = (
                successful[
                    "Propagation_Penalty"
                ]
                > tolerance
            )

            n_propagation = int(
                has_propagation.sum()
            )

            propagation_prevalence = (
                100.0
                * n_propagation
                / n_successful
            )

            # =================================================
            # AR target failure
            # =================================================

            if method == "ProACTS":

                # ProACTS success is already defined using
                # AR_Flip_Loss <= tolerance.
                #
                # Hence AR target failure among successful
                # ProACTS CFs is zero by construction.
                n_ar_failure = 0
                ar_failure_rate = 0.0

            else:

                # Baseline appeared successful under NA
                # evaluation. Check whether target success
                # survives autoregressive rollout.
                ar_target_failure = (
                    successful[
                        "AR_Flip_Loss"
                    ]
                    > tolerance
                )

                n_ar_failure = int(
                    ar_target_failure.sum()
                )

                ar_failure_rate = (
                    100.0
                    * n_ar_failure
                    / n_successful
                )

            # =================================================
            # Propagation severity
            # =================================================

            propagation_values = (
                successful[
                    "Propagation_Penalty"
                ]
                .to_numpy(
                    dtype=float
                )
            )

            (
                penalty_median,
                penalty_q1,
                penalty_q3,
            ) = _median_iqr_values(
                propagation_values
            )

            # =================================================
            # Dataset-method result
            # =================================================

            rows.append(
                {
                    "Dataset":
                        dataset,

                    "Method":
                        method,

                    "Success Semantics":
                        success_semantics,

                    "N Factual Queries":
                        n_queries,

                    "Target-Successful CFs":
                        n_successful,

                    "Target Success (%)":
                        (
                            100.0
                            * n_successful
                            / n_queries
                            if n_queries > 0
                            else np.nan
                        ),

                    "CFs with Propagation":
                        n_propagation,

                    "Propagation Prevalence (%)":
                        propagation_prevalence,

                    "AR Target Failures":
                        n_ar_failure,

                    "AR Target Failure (%)":
                        ar_failure_rate,

                    "Median Propagation Penalty":
                        penalty_median,

                    "Propagation Penalty [Q1, Q3]":
                        (
                            f"{_format_number(penalty_median)} "
                            f"[{_format_number(penalty_q1)}, "
                            f"{_format_number(penalty_q3)}]"
                        ),
                }
            )

    # ========================================================
    # Dataset-level table
    # ========================================================

    table = pd.DataFrame(
        rows
    )

    # ========================================================
    # Overall pooled rows
    # ========================================================

    overall_rows = []

    for method in methods:

        method_table = (
            table.loc[
                table["Method"]
                == method
            ]
            .copy()
        )

        # ----------------------------------------------------
        # Pooled counts
        # ----------------------------------------------------

        total_queries = int(
            method_table[
                "N Factual Queries"
            ].sum()
        )

        total_successful = int(
            method_table[
                "Target-Successful CFs"
            ].sum()
        )

        total_propagation = int(
            method_table[
                "CFs with Propagation"
            ].sum()
        )

        # ----------------------------------------------------
        # Pooled target success
        # ----------------------------------------------------

        overall_target_success = (
            100.0
            * total_successful
            / total_queries
            if total_queries > 0
            else np.nan
        )

        # ----------------------------------------------------
        # Pooled propagation prevalence
        # ----------------------------------------------------

        overall_propagation = (
            100.0
            * total_propagation
            / total_successful
            if total_successful > 0
            else np.nan
        )

        # ----------------------------------------------------
        # Pooled AR target failure
        # ----------------------------------------------------

        if method == "ProACTS":

            total_ar_failures = 0

            overall_ar_failure = (
                0.0
                if total_successful > 0
                else np.nan
            )

        else:

            total_ar_failures = int(
                method_table[
                    "AR Target Failures"
                ].sum()
            )

            overall_ar_failure = (
                100.0
                * total_ar_failures
                / total_successful
                if total_successful > 0
                else np.nan
            )

        # ----------------------------------------------------
        # Overall row
        # ----------------------------------------------------

        overall_rows.append(
            {
                "Dataset":
                    "Overall",

                "Method":
                    method,

                "Success Semantics":
                    (
                        "AR"
                        if method == "ProACTS"
                        else "NA"
                    ),

                "N Factual Queries":
                    total_queries,

                "Target-Successful CFs":
                    total_successful,

                "Target Success (%)":
                    overall_target_success,

                "CFs with Propagation":
                    total_propagation,

                "Propagation Prevalence (%)":
                    overall_propagation,

                "AR Target Failures":
                    total_ar_failures,

                "AR Target Failure (%)":
                    overall_ar_failure,

                "Median Propagation Penalty":
                    np.nan,

                "Propagation Penalty [Q1, Q3]":
                    "NA",
            }
        )

    # ========================================================
    # Add overall rows
    # ========================================================

    table = pd.concat(
        [
            table,
            pd.DataFrame(
                overall_rows
            ),
        ],
        ignore_index=True,
    )

    # ========================================================
    # Save Excel
    # ========================================================

    if output_path is not None:

        directory = os.path.dirname(
            output_path
        )

        if directory:

            os.makedirs(
                directory,
                exist_ok=True,
            )

        with pd.ExcelWriter(
            output_path,
            engine="openpyxl",
        ) as writer:

            table.to_excel(
                writer,
                sheet_name=(
                    "Propagation_Prevalence"
                ),
                index=False,
            )

        _autosize_workbook(
            output_path
        )

    return table



# ============================================================
# ProACTS propagation-loss ablation:
# full ProACTS vs ProACTS without propagation loss
# ============================================================


ABLATION_FULL_NAME = "ProACTS"
ABLATION_NO_PROP_NAME = "ProACTS w/o Prop."


ABLATION_METRICS = {
    "Propagation Penalty": "process_violation",
    "Margin Loss": "margin_loss",
    "Distance": "distance",
    "Sparsity": "sparsity",
}


# ============================================================
# Load ablated results
# ============================================================

def _load_proacts_ablation_results(
    all_dataset_results: Mapping,
    *,
    ablated_root: str,
    relative_path_template: str,
    sheet_name: str = "CF_Fitness_Analysis",
    fitness_column: str = "fitness",
) -> dict:
    """
    Load ProACTS-without-propagation-loss results.

    Parameters
    ----------
    all_dataset_results:
        Existing main results dictionary:

            all_dataset_results[dataset][seed]["ProACTS"]

    ablated_root:
        Root directory containing ablation results.

    relative_path_template:
        Relative path from ablated_root to each Excel file.

        It may contain:

            {dataset}
            {seed}

        Example:

            "{dataset}/{seed}/results.xlsx"

        so the final path becomes:

            <ablated_root>/<dataset>/<seed>/results.xlsx

    Returns
    -------
    dict

        result[dataset][seed] = DataFrame

    Notes
    -----
    One best counterfactual per factual query is selected
    using the same _load_best_counterfactuals() function as
    the main evaluation.
    """

    loaded = {}

    for (
        dataset,
        seed_results,
    ) in all_dataset_results.items():

        loaded[
            dataset
        ] = {}

        for seed in seed_results.keys():

            relative_path = (
                relative_path_template.format(
                    dataset=dataset,
                    seed=seed,
                )
            )

            excel_path = os.path.join(
                ablated_root,
                relative_path,
            )

            if not os.path.isfile(
                excel_path
            ):

                raise FileNotFoundError(
                    "Ablated result file not found:\n"
                    f"{excel_path}\n\n"
                    f"Dataset: {dataset}\n"
                    f"Seed: {seed}"
                )

            ablated_df = (
                _load_best_counterfactuals(
                    excel_path,
                    sheet_name=sheet_name,
                    fitness_column=fitness_column,
                )
            )

            loaded[
                dataset
            ][
                seed
            ] = ablated_df

    return loaded


# ============================================================
# Combine full and ablated ProACTS across seeds
# ============================================================

def _combine_proacts_ablation_across_seeds(
    full_seed_results: Mapping,
    ablated_seed_results: Mapping,
) -> pd.DataFrame:
    """
    Combine full ProACTS and ablated ProACTS across seeds.

    Output Method values:

        ProACTS
        ProACTS w/o Prop.
    """

    frames = []

    full_seeds = set(
        full_seed_results.keys()
    )

    ablated_seeds = set(
        ablated_seed_results.keys()
    )

    if (
        full_seeds
        != ablated_seeds
    ):

        raise ValueError(
            "Seed mismatch between full and "
            "ablated ProACTS results.\n"
            f"Full seeds: {sorted(full_seeds)}\n"
            f"Ablated seeds: {sorted(ablated_seeds)}"
        )

    for seed in sorted(
        full_seeds
    ):

        if "ProACTS" not in (
            full_seed_results[
                seed
            ]
        ):

            raise KeyError(
                f"Seed {seed} does not contain "
                "'ProACTS'."
            )

        # ----------------------------------------------------
        # Full ProACTS
        # ----------------------------------------------------

        full = (
            full_seed_results[
                seed
            ][
                "ProACTS"
            ]
            .copy()
        )

        full.insert(
            0,
            "Seed",
            seed,
        )

        full.insert(
            1,
            "Method",
            ABLATION_FULL_NAME,
        )

        frames.append(
            full
        )

        # ----------------------------------------------------
        # Ablated ProACTS
        # ----------------------------------------------------

        ablated = (
            ablated_seed_results[
                seed
            ]
            .copy()
        )

        ablated.insert(
            0,
            "Seed",
            seed,
        )

        ablated.insert(
            1,
            "Method",
            ABLATION_NO_PROP_NAME,
        )

        frames.append(
            ablated
        )

    if not frames:

        raise ValueError(
            "No ablation observations supplied."
        )

    return pd.concat(
        frames,
        ignore_index=True,
        sort=False,
    )


# ============================================================
# Ablation flip-success significance
# ============================================================

def _build_ablation_success_table(
    combined: pd.DataFrame,
    *,
    tolerance: float = 1e-8,
) -> pd.DataFrame:
    """
    Compare AR flip success between:

        ProACTS
        ProACTS w/o propagation loss

    Experimental unit:
        one factual query.

    Each query is first aggregated across seeds using median
    margin_loss_flipped.

    Statistical test:
        exact paired McNemar.
    """

    paired = _pair_metric(
        combined=combined,

        method_a=(
            ABLATION_FULL_NAME
        ),

        method_b=(
            ABLATION_NO_PROP_NAME
        ),

        column_a=(
            "margin_loss_flipped"
        ),

        column_b=(
            "margin_loss_flipped"
        ),
    )

    if paired.empty:

        return pd.DataFrame(
            [
                {
                    "N Paired": 0,

                    "ProACTS Successful":
                        0,

                    "Ablated Successful":
                        0,

                    "ProACTS Success %":
                        np.nan,

                    "Ablated Success %":
                        np.nan,

                    "Both Successful":
                        0,

                    "ProACTS Only":
                        0,

                    "Ablated Only":
                        0,

                    "Neither":
                        0,

                    "Discordant":
                        0,

                    "McNemar Raw p":
                        np.nan,

                    "Direction":
                        "Undefined",
                }
            ]
        )

    # ========================================================
    # Binary success
    # ========================================================

    full_success = (
        paired[
            "A"
        ]
        .to_numpy(
            dtype=float
        )
        <= tolerance
    )

    ablated_success = (
        paired[
            "B"
        ]
        .to_numpy(
            dtype=float
        )
        <= tolerance
    )

    n = int(
        len(
            paired
        )
    )

    # ========================================================
    # Contingency counts
    # ========================================================

    both = int(
        (
            full_success
            &
            ablated_success
        ).sum()
    )

    full_only = int(
        (
            full_success
            &
            ~ablated_success
        ).sum()
    )

    ablated_only = int(
        (
            ~full_success
            &
            ablated_success
        ).sum()
    )

    neither = int(
        (
            ~full_success
            &
            ~ablated_success
        ).sum()
    )

    discordant = (
        full_only
        +
        ablated_only
    )

    # ========================================================
    # Exact paired McNemar
    # ========================================================

    if discordant == 0:

        raw_p = 1.0

    else:

        result = mcnemar(
            [
                [
                    both,
                    full_only,
                ],
                [
                    ablated_only,
                    neither,
                ],
            ],
            exact=True,
        )

        raw_p = float(
            result.pvalue
        )

    # ========================================================
    # Direction
    # ========================================================

    if full_only > ablated_only:

        direction = (
            ABLATION_FULL_NAME
        )

    elif ablated_only > full_only:

        direction = (
            ABLATION_NO_PROP_NAME
        )

    else:

        direction = "Tie"

    return pd.DataFrame(
        [
            {
                "N Paired":
                    n,

                "ProACTS Successful":
                    int(
                        full_success.sum()
                    ),

                "Ablated Successful":
                    int(
                        ablated_success.sum()
                    ),

                "ProACTS Success %":
                    (
                        100.0
                        *
                        full_success.mean()
                    ),

                "Ablated Success %":
                    (
                        100.0
                        *
                        ablated_success.mean()
                    ),

                "Both Successful":
                    both,

                "ProACTS Only":
                    full_only,

                "Ablated Only":
                    ablated_only,

                "Neither":
                    neither,

                "Discordant":
                    discordant,

                "McNemar Raw p":
                    raw_p,

                "Direction":
                    direction,
            }
        ]
    )


# ============================================================
# Ablation continuous statistical table
# ============================================================

def _build_ablation_conditional_table(
    combined: pd.DataFrame,
    *,
    tolerance: float = 1e-8,
) -> pd.DataFrame:
    """
    Compare continuous metrics between full and ablated
    ProACTS.

    Only factual queries where BOTH variants achieve the
    desired target under autoregressive evaluation are kept.

    Success:
        median margin_loss_flipped <= tolerance

    Metrics:
        Propagation Penalty
        Margin Loss
        Distance
        Sparsity

    Statistical tests:
        Wilcoxon
        Pratt
        Z-split

    Effect:
        rank-biserial correlation on non-tied pairs.
    """

    # ========================================================
    # Determine common AR-successful queries
    # ========================================================

    success_pairs = _pair_metric(
        combined=combined,

        method_a=(
            ABLATION_FULL_NAME
        ),

        method_b=(
            ABLATION_NO_PROP_NAME
        ),

        column_a=(
            "margin_loss_flipped"
        ),

        column_b=(
            "margin_loss_flipped"
        ),
    )

    success_a = (
        success_pairs[
            "A"
        ]
        .to_numpy(
            dtype=float
        )
        <= tolerance
    )

    success_b = (
        success_pairs[
            "B"
        ]
        .to_numpy(
            dtype=float
        )
        <= tolerance
    )

    common_success = (
        success_a
        &
        success_b
    )

    retained_ids = (
        success_pairs.loc[
            common_success,
            ID_COLUMNS,
        ]
        .drop_duplicates()
        .reset_index(
            drop=True
        )
    )

    n_before = int(
        len(
            success_pairs
        )
    )

    rows = []

    # ========================================================
    # Metrics
    # ========================================================

    for (
        metric_name,
        metric_column,
    ) in ABLATION_METRICS.items():

        paired_metric = _pair_metric(
            combined=combined,

            method_a=(
                ABLATION_FULL_NAME
            ),

            method_b=(
                ABLATION_NO_PROP_NAME
            ),

            column_a=(
                metric_column
            ),

            column_b=(
                metric_column
            ),
        )

        conditional = (
            retained_ids.merge(
                paired_metric,
                on=ID_COLUMNS,
                how="inner",
                validate="one_to_one",
            )
        )

        # ----------------------------------------------------
        # Nothing retained
        # ----------------------------------------------------

        if conditional.empty:

            rows.append(
                {
                    "Metric":
                        metric_name,

                    "N Paired Before Condition":
                        n_before,

                    "N Retained":
                        0,

                    "Retained %":
                        0.0,

                    "N Non-Ties":
                        0,

                    "Tie %":
                        np.nan,

                    "ProACTS Mean ± SD":
                        "NA",

                    "Ablated Mean ± SD":
                        "NA",

                    "ProACTS Median [Q1, Q3]":
                        "NA",

                    "Ablated Median [Q1, Q3]":
                        "NA",

                    "Median Difference "
                    "ProACTS-Ablated":
                        np.nan,

                    "Wilcox Raw p":
                        np.nan,

                    "Pratt Raw p":
                        np.nan,

                    "Z-split Raw p":
                        np.nan,

                    "Rank-Biserial (Non-Ties)":
                        np.nan,

                    "Effect":
                        "Undefined",

                    "Better Method":
                        "Undefined",
                }
            )

            continue

        # ====================================================
        # Values
        # ====================================================

        full_values = (
            conditional[
                "A"
            ]
            .to_numpy(
                dtype=float
            )
        )

        ablated_values = (
            conditional[
                "B"
            ]
            .to_numpy(
                dtype=float
            )
        )

        difference = (
            full_values
            -
            ablated_values
        )

        ties = np.isclose(
            difference,
            0.0,
            atol=tolerance,
            rtol=0.0,
        )

        # ====================================================
        # Signed-rank tests
        # ====================================================

        (
            wilcox_stat,
            wilcox_p,
        ) = _run_wilcoxon(
            full_values,
            ablated_values,
            zero_tolerance=tolerance,
        )

        (
            pratt_stat,
            pratt_p,
        ) = _run_pratt(
            full_values,
            ablated_values,
            zero_tolerance=tolerance,
        )

        (
            zsplit_stat,
            zsplit_p,
        ) = _run_zsplit(
            full_values,
            ablated_values,
            zero_tolerance=tolerance,
        )

        # ====================================================
        # Effect size
        # ====================================================

        effect = (
            _rank_biserial_correlation(
                full_values,
                ablated_values,
                zero_tolerance=tolerance,
            )
        )

        # All selected metrics are lower-is-better.
        better = (
            _better_method_from_effect(
                ABLATION_FULL_NAME,
                ABLATION_NO_PROP_NAME,
                effect,
                lower_is_better=True,
                tolerance=tolerance,
            )
        )

        # ====================================================
        # Result
        # ====================================================

        rows.append(
            {
                "Metric":
                    metric_name,

                "N Paired Before Condition":
                    n_before,

                "N Retained":
                    int(
                        len(
                            conditional
                        )
                    ),

                "Retained %":
                    (
                        100.0
                        *
                        len(
                            conditional
                        )
                        /
                        n_before
                        if n_before > 0
                        else np.nan
                    ),

                "N Non-Ties":
                    int(
                        (
                            ~ties
                        ).sum()
                    ),

                "Tie %":
                    (
                        100.0
                        *
                        ties.mean()
                    ),

                "ProACTS Mean ± SD":
                    _mean_sd_text(
                        full_values
                    ),

                "Ablated Mean ± SD":
                    _mean_sd_text(
                        ablated_values
                    ),

                "ProACTS Median [Q1, Q3]":
                    _median_iqr_text(
                        full_values
                    ),

                "Ablated Median [Q1, Q3]":
                    _median_iqr_text(
                        ablated_values
                    ),

                "Median Difference "
                "ProACTS-Ablated":
                    float(
                        np.median(
                            difference
                        )
                    ),

                "Wilcox Statistic":
                    wilcox_stat,

                "Wilcox Raw p":
                    wilcox_p,

                "Pratt Statistic":
                    pratt_stat,

                "Pratt Raw p":
                    pratt_p,

                "Z-split Statistic":
                    zsplit_stat,

                "Z-split Raw p":
                    zsplit_p,

                "Rank-Biserial (Non-Ties)":
                    effect,

                "Effect":
                    _effect_magnitude(
                        effect
                    ),

                "Better Method":
                    better,
            }
        )

    return pd.DataFrame(
        rows
    )


# ============================================================
# Full ablation analysis across all datasets
# ============================================================

def summarize_proacts_propagation_ablation(
    all_dataset_results: Mapping,
    *,
    ablated_root: str,
    relative_path_template: str,
    output_path: Optional[str] = (
        "summary_evaluation/"
        "proacts_propagation_ablation.xlsx"
    ),
    alpha: float = 0.05,
    tolerance: float = 1e-8,
    practical_effect_threshold: float = 0.10,
    sheet_name: str = "CF_Fitness_Analysis",
    fitness_column: str = "fitness",
) -> dict:
    """
    Statistical ablation analysis:

        full ProACTS
            vs
        ProACTS without propagation loss.

    Produces two statistical tables:

        1. Flip_Success
           Exact paired McNemar test.

        2. Conditional_Metrics
           Continuous metrics on queries where both variants
           achieve AR flip success.

    Multiple-testing correction
    ---------------------------
    Success:
        Holm across datasets.

    Continuous metrics:
        Holm across datasets separately for each metric
        and each zero-handling convention.

    Experimental unit:
        one factual query after median aggregation across seeds.
    """

    # ========================================================
    # Load ablated experiment files
    # ========================================================

    ablated_results = (
        _load_proacts_ablation_results(
            all_dataset_results,

            ablated_root=(
                ablated_root
            ),

            relative_path_template=(
                relative_path_template
            ),

            sheet_name=(
                sheet_name
            ),

            fitness_column=(
                fitness_column
            ),
        )
    )

    success_frames = []
    conditional_frames = []

    # ========================================================
    # Dataset analysis
    # ========================================================

    for (
        dataset,
        full_seed_results,
    ) in all_dataset_results.items():

        combined = (
            _combine_proacts_ablation_across_seeds(
                full_seed_results=(
                    full_seed_results
                ),

                ablated_seed_results=(
                    ablated_results[
                        dataset
                    ]
                ),
            )
        )

        # ----------------------------------------------------
        # Success
        # ----------------------------------------------------

        success = (
            _build_ablation_success_table(
                combined,
                tolerance=tolerance,
            )
        )

        success.insert(
            0,
            "Dataset",
            dataset,
        )

        success_frames.append(
            success
        )

        # ----------------------------------------------------
        # Conditional continuous metrics
        # ----------------------------------------------------

        conditional = (
            _build_ablation_conditional_table(
                combined,
                tolerance=tolerance,
            )
        )

        conditional.insert(
            0,
            "Dataset",
            dataset,
        )

        conditional_frames.append(
            conditional
        )

    # ========================================================
    # Combine datasets
    # ========================================================

    success_table = pd.concat(
        success_frames,
        ignore_index=True,
    )

    conditional_table = pd.concat(
        conditional_frames,
        ignore_index=True,
    )

    # ========================================================
    # Holm correction: McNemar across datasets
    # ========================================================

    success_table[
        "McNemar Holm-adjusted p"
    ] = np.nan

    valid = success_table.index[
        np.isfinite(
            success_table[
                "McNemar Raw p"
            ]
        )
    ]

    if len(
        valid
    ) > 0:

        (
            _,
            adjusted,
            _,
            _,
        ) = multipletests(
            success_table.loc[
                valid,
                "McNemar Raw p",
            ].to_numpy(
                dtype=float
            ),

            alpha=alpha,

            method="holm",
        )

        success_table.loc[
            valid,
            "McNemar Holm-adjusted p",
        ] = adjusted

    success_table[
        "Holm p"
    ] = (
        success_table[
            "McNemar Holm-adjusted p"
        ]
        .apply(
            _format_p_value
        )
    )

    # ========================================================
    # Success conclusion
    # ========================================================

    success_conclusions = []

    for _, row in (
        success_table.iterrows()
    ):

        p = row[
            "McNemar Holm-adjusted p"
        ]

        direction = row[
            "Direction"
        ]

        if (
            np.isfinite(p)
            and
            p < alpha
            and
            direction != "Tie"
        ):

            success_conclusions.append(
                f"{direction} higher success"
            )

        else:

            success_conclusions.append(
                "No significant difference"
            )

    success_table[
        "Conclusion"
    ] = (
        success_conclusions
    )

    # ========================================================
    # Holm correction:
    # continuous metrics across datasets
    # ========================================================

    for adjusted_column in [
        "Wilcox Holm-adjusted p",
        "Pratt Holm-adjusted p",
        "Z-split Holm-adjusted p",
    ]:

        conditional_table[
            adjusted_column
        ] = np.nan

    # --------------------------------------------------------
    # Each metric is a separate hypothesis family.
    #
    # Within a metric:
    #     one test per dataset.
    #
    # Wilcox / Pratt / Z-split remain separate sensitivity
    # analyses of the same hypotheses.
    # --------------------------------------------------------

    for (
        metric,
        group,
    ) in conditional_table.groupby(
        "Metric",
        sort=False,
    ):

        for (
            raw_column,
            adjusted_column,
        ) in [
            (
                "Wilcox Raw p",
                "Wilcox Holm-adjusted p",
            ),
            (
                "Pratt Raw p",
                "Pratt Holm-adjusted p",
            ),
            (
                "Z-split Raw p",
                "Z-split Holm-adjusted p",
            ),
        ]:

            valid = group.index[
                np.isfinite(
                    conditional_table.loc[
                        group.index,
                        raw_column,
                    ]
                )
            ]

            if len(
                valid
            ) == 0:

                continue

            (
                _,
                adjusted,
                _,
                _,
            ) = multipletests(
                conditional_table.loc[
                    valid,
                    raw_column,
                ].to_numpy(
                    dtype=float
                ),

                alpha=alpha,

                method="holm",
            )

            conditional_table.loc[
                valid,
                adjusted_column,
            ] = adjusted

    # ========================================================
    # Formatted p-values
    # ========================================================

    conditional_table[
        "Wilcox Holm p"
    ] = (
        conditional_table[
            "Wilcox Holm-adjusted p"
        ]
        .apply(
            _format_p_value
        )
    )

    conditional_table[
        "Pratt Holm p"
    ] = (
        conditional_table[
            "Pratt Holm-adjusted p"
        ]
        .apply(
            _format_p_value
        )
    )

    conditional_table[
        "Z-split Holm p"
    ] = (
        conditional_table[
            "Z-split Holm-adjusted p"
        ]
        .apply(
            _format_p_value
        )
    )

    # ========================================================
    # Robustness + conclusion
    # ========================================================

    robustness = []
    conclusions = []

    for _, row in (
        conditional_table.iterrows()
    ):

        wilcox_p = row[
            "Wilcox Holm-adjusted p"
        ]

        pratt_p = row[
            "Pratt Holm-adjusted p"
        ]

        zsplit_p = row[
            "Z-split Holm-adjusted p"
        ]

        effect = row[
            "Rank-Biserial (Non-Ties)"
        ]

        better = row[
            "Better Method"
        ]

        # ----------------------------------------------------
        # Statistical significance
        # ----------------------------------------------------

        wilcox_sig = (
            np.isfinite(
                wilcox_p
            )
            and
            wilcox_p < alpha
        )

        pratt_sig = (
            np.isfinite(
                pratt_p
            )
            and
            pratt_p < alpha
        )

        zsplit_sig = (
            np.isfinite(
                zsplit_p
            )
            and
            zsplit_p < alpha
        )

        significance = [
            wilcox_sig,
            pratt_sig,
            zsplit_sig,
        ]

        # ----------------------------------------------------
        # Robustness
        # ----------------------------------------------------

        if all(
            significance
        ):

            robustness.append(
                "Robust: all significant"
            )

        elif not any(
            significance
        ):

            robustness.append(
                "Robust: none significant"
            )

        else:

            robustness.append(
                "Sensitive to zero handling"
            )

        # ----------------------------------------------------
        # Practical effect
        # ----------------------------------------------------

        practical = (
            np.isfinite(
                effect
            )
            and
            abs(
                effect
            )
            >=
            practical_effect_threshold
        )

        # ----------------------------------------------------
        # Primary conclusion:
        # Wilcox remains primary.
        # ----------------------------------------------------

        if (
            wilcox_sig
            and
            practical
            and
            better
            not in {
                "Tie",
                "Undefined",
            }
        ):

            conclusion = (
                f"{better} better"
            )

        elif (
            wilcox_sig
            and
            not practical
        ):

            conclusion = (
                "Statistically significant; "
                "negligible effect"
            )

        elif (
            not wilcox_sig
            and
            practical
        ):

            conclusion = (
                "Non-negligible effect; "
                "not statistically significant"
            )

        else:

            conclusion = (
                "No meaningful evidence "
                "of difference"
            )

        conclusions.append(
            conclusion
        )

    conditional_table[
        "Zero-Handling Robustness"
    ] = robustness

    conditional_table[
        "Conclusion"
    ] = conclusions

    # ========================================================
    # Final columns
    # ========================================================

    success_table = success_table[
        [
            "Dataset",
            "N Paired",

            "ProACTS Successful",
            "Ablated Successful",

            "ProACTS Success %",
            "Ablated Success %",

            "Both Successful",
            "ProACTS Only",
            "Ablated Only",
            "Neither",
            "Discordant",

            "McNemar Raw p",
            "McNemar Holm-adjusted p",
            "Holm p",

            "Direction",
            "Conclusion",
        ]
    ].copy()

    conditional_table = conditional_table[
        [
            "Dataset",
            "Metric",

            "N Paired Before Condition",
            "N Retained",
            "Retained %",

            "N Non-Ties",
            "Tie %",

            "ProACTS Mean ± SD",
            "Ablated Mean ± SD",

            "ProACTS Median [Q1, Q3]",
            "Ablated Median [Q1, Q3]",

            "Median Difference ProACTS-Ablated",

            "Wilcox Raw p",
            "Wilcox Holm-adjusted p",
            "Wilcox Holm p",

            "Pratt Holm p",
            "Z-split Holm p",

            "Rank-Biserial (Non-Ties)",
            "Effect",

            "Better Method",

            "Zero-Handling Robustness",
            "Conclusion",
        ]
    ].copy()

    # ========================================================
    # Save statistical tables only
    # ========================================================

    if output_path is not None:

        directory = os.path.dirname(
            output_path
        )

        if directory:

            os.makedirs(
                directory,
                exist_ok=True,
            )

        with pd.ExcelWriter(
            output_path,
            engine="openpyxl",
        ) as writer:

            success_table.to_excel(
                writer,
                sheet_name="Flip_Success",
                index=False,
            )

            conditional_table.to_excel(
                writer,
                sheet_name="Conditional_Metrics",
                index=False,
            )

        _autosize_workbook(
            output_path
        )

    return {
        "success": success_table,
        "conditional": conditional_table,
    }
